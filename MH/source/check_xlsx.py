# Cross-check the recalculated xlsx against model.json (run after scripts/recalc.py).
import os, sys, json
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import model as MD
M = json.load(open(MD.OUT, encoding='utf-8'))
R = json.load(open(os.path.join(HERE, '_xlsx_rows.json')))
wb = load_workbook(os.path.join(HERE, '..', 'MH_Robotics_Financial_Model.xlsx'), data_only=True)
FCOLS = ['C', 'D', 'E', 'F', 'G']
bad = 0; n = 0

def cmp(label, x, y, tol=0.01):
    global bad, n
    n += 1
    if x is None or abs((x or 0) - y) > tol * max(1, abs(y)) * 0.01 + 1e-6:
        bad += 1; print('MISMATCH', label, x, y)

for s in MD.SC:
    ws = wb[f'FM_{MD.SCN[s]}']; L = M['scenarios'][s]
    for name, row in R['fm_rows'][s].items():
        if name in L and isinstance(L[name], list):
            for t in range(5):
                cmp(f'{s}.{name}.Y{t+1}', ws[f'{FCOLS[t]}{row}'].value, L[name][t])
ws = wb['Household']; hr = R['hh_rows']
for col, key in (('B', 'purchase_direct_Y3'), ('C', 'purchase_direct_Y5'), ('D', 'rental_direct_Y3'), ('E', 'rental_direct_Y5'),
                 ('F', 'retrofit_purchase_Y3'), ('G', 'newbuild_purchase_Y3')):
    h = M['household'][key]
    cmp(key + '.rev5', ws[f"{col}{hr['5년 매출']}"].value, h['rev5'])
    cmp(key + '.contrib5', ws[f"{col}{hr['5년 Lifetime Contribution']}"].value, h['contrib5'])
    cmp(key + '.gp5', ws[f"{col}{hr['5년 매출총이익 (채널비용 전)']}"].value, h['gp5'])
    cmp(key + '.y0', ws[f"{col}{hr['Year 0 매출 (INSTALL)']}"].value, h['y0'])
    cmp(key + '.operate', ws[f"{col}{hr['OPERATE 매출 (5년)']}"].value, h['layers']['operate'])
    cmp(key + '.service', ws[f"{col}{hr['5년 서비스 원가 (Care · Consumables · Warranty)']}"].value, h['service_cost5'])
ws = wb['Unit_Economics']; ur = R['ue_rows']
for col, lab in (('B', 'Y3'), ('C', 'Y5')):
    r = M['rental'][lab]
    cmp('rental.payback.' + lab, ws[f"{col}{ur['Payback (BOM ÷ (요금 − Care − Grip), 개월)']}"].value, r['payback'])
    cmp('rental.contrib.' + lab, ws[f"{col}{ur['월 Contribution']}"].value, r['contrib_m'])
    cmp('rental.bommax.' + lab, ws[f"{col}{ur['Payback 허들 충족 최대 BOM']}"].value, r['bom_max'])
ux = R['ue_extra']
cmp('partner_irr', ws[f"B{ux['irr_row']}"].value, M['partner_irr']['B']['irr_y'])
for col in 'BCD':
    cmp('bom_check_' + col, ws[f"{col}{ux['bom_check_row']}"].value + 1, 1)
ws = wb['Market']; mk = M['market']['B']; MR = R['mk_rows']
for lab, key in (('교차검증 ①: 20년+ 아파트(2023) ÷ 교체주기', 'tri1'), ('교차검증 ②: 매매거래 × 아파트 비중 × 교체율 + 비거래 교체', 'tri2'),
                 ('① Remodeling 대상 세대', 'fit'), ('① SAM Remodeling', 'sam_remodel'), ('Retrofit 호환 세대 Pool (재고)', 'retro_pool'),
                 ('② SAM Retrofit', 'sam_retro'), ('③ SAM New-build', 'sam_new'), ('구매 고객 ARPU (Care × 가입률 + Consumables × 구매율)', 'arpu'),
                 ('참고 TAM: Premium 세대 (Remodeling + New-build) × 전체 패키지', 'tam'), ('SAM 합계 (① + ② + ③)', 'sam'),
                 ('SOM: Y5 매출 (Base Plan)', 'som'), ('④ Y5 OPERATE 매출 (Base Plan)', 'recurring_y5')):
    cmp('market.' + key, ws[f'B{MR[lab]}'].value, mk[key])
ws = wb['Budget_24M']; BR = R['bud_rows']; F = M['funding']; TP = F['tips']
cmp('bud.people_y1', ws[f"K{BR['team_tot']}"].value, F['people'][0])
cmp('bud.people_y2', ws[f"L{BR['team_tot']}"].value, F['people'][1])
cmp('bud.fte_y1', ws[f"I{BR['fte']}"].value, F['fte'][0])
cmp('bud.fte_y2', ws[f"J{BR['fte']}"].value, F['fte'][1])
for col in 'KL':
    cmp('bud.people_chk_' + col, ws[f"{col}{BR['people_chk']}"].value + 1, 1)
for col in 'EF':
    cmp('bud.other_chk_' + col, ws[f"{col}{BR['other_chk']}"].value + 1, 1)
cmp('bud.tips_total', ws[f"G{BR['tot']}"].value, TP['total'])
cmp('bud.tips_indirect', ws[f"G{BR['ind']}"].value, TP['indirect'])
for i, d in enumerate(TP['rows'][:-1]):
    pass
cmp('bud.cash_ok', ws[f"G{BR['현금 비율 하한 충족 (1 = 예)']}"].value, 1)
cmp('bud.private_cash', ws[f"G{BR['기관부담 중 현금']}"].value, TP['private_cash'])
cmp('bud.indirect_rate', ws[f"G{BR['간접비율 (간접비 ÷ 현금 직접비)']}"].value, TP['indirect_rate'])
cmp('bud.spend', ws[f"G{BR['spend']}"].value, F['spend_total'])
cmp('bud.gov', ws[f"G{BR['gov']}"].value, TP['gov'])
cmp('bud.buffer', ws[f"G{BR['buf']}"].value, F['buffer'])
cmp('bud.seed_base', ws[f"G{BR['seed']}"].value, F['seed_base'])
ws = wb['Sensitivity']
cmp('sens.base', ws['M5'].value, M['sens_household']['base'])
print(f'checked {n} values, mismatches {bad}')
sys.exit(1 if bad else 0)
