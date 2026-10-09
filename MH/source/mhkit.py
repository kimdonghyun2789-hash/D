# Shared idioms for the MH Robotics main deck (v5). Light theme, two weights, orange only for Robot Zone / Path / Key Number.
import os, json
from PIL import Image
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE, MSO_PATTERN_TYPE
import kit
from kit import T, W, H, MX, CW, text, rect, hline, vline, table, chip, arrow, dashed_rect, alpha, text_h, footer, dot
import common
from common import M, META, start

RD = os.path.join(common.ROOT, 'assets', 'renders')
GREY = '8C9198'
EDGE = 'C9CDD2'
INK = '15171A'
INK2 = '4A4F57'
SOFT = 'F4F5F6'
SOFT2 = 'ECEEF0'
ACC = 'E2571B'
ACC_SOFT = 'FBEDE6'
FOOT_LEFT = 'MH Robotics  ·  Seed · TIPS IR  ·  2026.10'
MT_LABEL = {'TBV': 'TO BE VALIDATED', 'FUTURE': 'FUTURE CONCEPT'}
PAGE = {'n': 0}


def mt(s, x, y, kind, size=6.5, h=0.17, label=None, fill=None):
    """Low-priority Number Tag: outlined, grey text, no fill (fill only when placed over an image)."""
    lab = label or MT_LABEL.get(kind, kind)
    w = kit.text_w(lab, size, True) + 0.12
    rect(s, x, y, w, h, fill=fill, line=EDGE, lw=0.5)
    kit.NOLOG['on'] = True
    text(s, x, y, w, h, lab, size=size, bold=True, color=GREY, align='c', anchor='m', check=False)
    kit.NOLOG['on'] = False
    kit._log(f'[{lab}]')
    return w


def mts(s, x, y, kinds, gap=0.05, **kw):
    for k in kinds:
        x += mt(s, x, y, k, **kw) + gap
    return x


def pg(prs):
    PAGE['n'] = len(prs.slides) + 1
    return PAGE['n']


def mfoot(s, note=None):
    footer(s, PAGE['n'], left=FOOT_LEFT, note=note)


def mhead(s, kicker, title, sub=None, size=25):
    text(s, MX, 0.46, 10.5, 0.24, kicker, size=10, bold=True, color=GREY, label='kicker')
    th = kit.text_h(title, size, CW, line=0.95, bold=True)
    text(s, MX, 0.74, CW, th + 0.04, title, size=size, bold=True, line=0.95, label='title:' + title[:12])
    y = 0.74 + th + 0.1
    if sub:
        sh = kit.text_h(sub, 12.5, CW)
        text(s, MX, y, CW, sh + 0.02, sub, size=12.5, color=INK2, label='sub:' + sub[:12])
        y += sh + 0.04
    return y + 0.14


def seg(s, x1, y1, x2, y2, color=None, lw=0.75, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color or INK2); c.line.width = Pt(lw)
    kit._nostyle(c)
    if dash: c.line.dash_style = MSO_LINE.DASH
    return c


def render(s, name, x, y, w, h, focus=(0.5, 0.5), zoom=1.0, bg=(255, 255, 255), border=False, src=None):
    """Cover-fit a render into the box (same crop as kit.image). Returns at(key | (px, py)) -> slide inches."""
    path = src or os.path.join(RD, name + '.png')
    iw, ih = Image.open(path).size; ar = w / h
    cw, ch = (int(ih * ar), ih) if iw / ih > ar else (iw, int(iw / ar))
    cw, ch = int(cw / zoom), int(ch / zoom)
    # same crop as kit.image (zoom < 1 pads the image centred on the background colour)
    cx = -((cw - iw) // 2) if cw > iw else min(max(int(focus[0] * iw - cw / 2), 0), iw - cw)
    cy = -((ch - ih) // 2) if ch > ih else min(max(int(focus[1] * ih - ch / 2), 0), ih - ch)
    kit.image(s, path, x, y, w, h, focus=focus, bg=bg, zoom=zoom)
    if border: rect(s, x, y, w, h, line=T['line'], lw=0.75)
    jp = os.path.join(RD, name + '.json')
    A = json.load(open(jp, encoding='utf-8')).get('anchors', {}) if os.path.exists(jp) else {}
    def at(key, dx=0.0, dy=0.0):
        px, py = A[key] if isinstance(key, str) else key
        return x + (px - cx) / cw * w + dx, y + (py - cy) / ch * h + dy
    return at


def faded(name, left=0.0, bottom=0.0, bg=(255, 255, 255)):
    """Render composited on bg with a soft fade to bg along the left / bottom edges (for full-bleed placement)."""
    im = Image.open(os.path.join(RD, name + '.png')).convert('RGBA')
    base = Image.new('RGBA', im.size, bg + (255,)); base.alpha_composite(im)
    iw, ih = im.size
    mask = Image.new('L', im.size, 255); px = mask.load()
    lw_, bh = int(iw * left), int(ih * bottom)
    for xx in range(lw_):
        a = int(255 * (xx / lw_) ** 1.6)
        for yy in range(ih): px[xx, yy] = min(px[xx, yy], a)
    for yy in range(ih - bh, ih):
        a = int(255 * ((ih - yy) / bh) ** 1.6)
        for xx in range(iw): px[xx, yy] = min(px[xx, yy], a)
    out = Image.new('RGBA', im.size, bg + (255,)); out.paste(base, (0, 0), mask)
    os.makedirs(kit.IMG_CACHE, exist_ok=True)
    p = os.path.join(kit.IMG_CACHE, f'{name}_fade_{bg[0]}.png'); out.convert('RGB').save(p)
    return p


def chipl(s, x, y, txt, size=9, side='r', fill='FFFFFF', color=None, line=None, bold=True):
    """Label chip vertically centred on y; side='r' extends right of x, 'l' left of x, 'c' centred."""
    w = kit.text_w(txt, size, bold) + 0.18; h = size * kit.LH / 72 + 0.08
    x0 = {'r': x, 'l': x - w, 'c': x - w / 2}[side]
    rect(s, x0, y - h / 2, w, h, fill=fill, line=line or EDGE, lw=0.5)
    text(s, x0, y - h / 2, w, h, txt, size=size, bold=bold, color=color or INK, align='c', anchor='m', check=False)
    return x0, w, h


def callout(s, at, key, txt, dx, dy, size=9, color=None, side=None):
    ax, ay = at(key); lx, ly = ax + dx, ay + dy
    seg(s, ax, ay, lx, ly, color='5C6169', lw=0.6)
    dot(s, ax, ay, 0.07, fill=INK)
    chipl(s, lx, ly, txt, size=size, side=side or ('r' if dx >= 0 else 'l'), color=color)


def marker(s, cx, cy, n, d=0.22, fill=None, color='FFFFFF', size=8.5, ring=True):
    if ring: rect(s, cx - d / 2 - 0.025, cy - d / 2 - 0.025, d + 0.05, d + 0.05, fill='FFFFFF', shape=MSO_SHAPE.OVAL)
    rect(s, cx - d / 2, cy - d / 2, d, d, fill=fill or INK, shape=MSO_SHAPE.OVAL)
    text(s, cx - d / 2, cy - d / 2, d, d, str(n), size=size, bold=True, color=color, align='c', anchor='m', check=False)


def knum(s, x, y, w, value, label, tag=None, vsize=28, color=None, lsize=10, tag_y=None):
    vh = vsize * kit.LH / 72
    text(s, x, y, w, vh, value, size=vsize, bold=True, color=color or ACC, label='knum ' + value)
    lh = kit.text_h(label, lsize, w)
    text(s, x, y + vh - 0.02, w, lh + 0.02, label, size=lsize, color=INK2, label='knuml ' + label[:10])
    yy = tag_y if tag_y is not None else y + vh + lh + 0.02
    if tag:
        mts(s, x, yy, [tag] if isinstance(tag, str) else tag); yy += 0.2
    return yy


def note(s, txt, y=None, size=9, color=None, w=None):
    """Grey source / caveat line just above the footer."""
    ww = w or CW
    hh = kit.text_h(txt, size, ww)
    yy = y if y is not None else H - 0.6 - hh
    text(s, MX, yy, ww, hh + 0.02, txt, size=size, color=color or GREY, label='note')
    return yy


def card(s, x, y, w, h, title, body=None, size=10, title_size=12.5, fill=None, line=None, title_color=None,
         body_color=None, bullet=None, kicker=None, kicker_color=None, pad=0.16, space_after=2):
    rect(s, x, y, w, h, fill=fill or SOFT, line=line)
    yy = y + pad - 0.02
    if kicker:
        text(s, x + pad, yy, w - 2 * pad, 0.22, kicker, size=8.5, bold=True, color=kicker_color or GREY, check=False)
        yy += 0.24
    th = kit.text_h(title, title_size, w - 2 * pad, bold=True)
    text(s, x + pad, yy, w - 2 * pad, th + 0.02, title, size=title_size, bold=True, color=title_color or INK,
         label='card:' + str(title)[:10])
    yy += th + 0.06
    if body:
        bh = y + h - yy - 0.08
        text(s, x + pad, yy, w - 2 * pad, bh, body, size=size, color=body_color or INK2, line=1.0, bullet=bullet,
             space_after=space_after, label='cardbody:' + str(title)[:10])
    return yy


def bar(s, x, y, w, h, txt, size=12.5, fill=None, color='FFFFFF', bold=True, align='c'):
    rect(s, x, y, w, h, fill=fill or INK)
    text(s, x + 0.2, y, w - 0.4, h, txt, size=size, bold=bold, color=color, align=align, anchor='m', label='bar')


def harrow(s, x1, y, x2, color=None, lw=1.5):
    arrow(s, x1, y, x2, y, color=color or GREY, lw=lw)


def varrow(s, x, y1, y2, color=None, lw=1.5):
    arrow(s, x, y1, x, y2, color=color or GREY, lw=lw)


def pill(s, x, y, txt, size=8.5, fill=None, color=None, bold=True, h=0.24, line=None):
    w = kit.text_w(txt, size, bold) + 0.22
    rect(s, x, y, w, h, fill=fill or SOFT2, line=line)
    text(s, x, y, w, h, txt, size=size, bold=bold, color=color or INK, align='c', anchor='m', check=False)
    return w


def eok(v, d=1):
    return f"{v / 10000:,.{d}f}억원"


def man(v, d=0):
    return f"{v:,.{d}f}만원"


def A(k):
    """Base input value."""
    for d in M['inputs']:
        if d['key'] == k: return d['vals']['B']
    raise KeyError(k)
