# Formula-driven Excel version of model.py (inputs in blue, formulas in black, cross-sheet links in green).
#   python3 MH/source/xlsx_model.py   -> MH/MH_Robotics_Financial_Model.xlsx  (then recalc + cross-check)
import os, sys, json, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import model as MD

ROOT = os.path.abspath(os.path.join(HERE, '..'))
OUT = os.path.join(ROOT, 'MH_Robotics_Financial_Model.xlsx')
FONT = 'Arial'
BLUE = Font(name=FONT, color='0000FF', size=10)
BLACK = Font(name=FONT, color='000000', size=10)
GREEN = Font(name=FONT, color='008000', size=10)
BOLD = Font(name=FONT, bold=True, size=10)
TITLE = Font(name=FONT, bold=True, size=14)
HDR = Font(name=FONT, bold=True, size=10, color='FFFFFF')
HDR_FILL = PatternFill('solid', fgColor='15171A')
SEC_FILL = PatternFill('solid', fgColor='F4F5F6')
KEY_FILL = PatternFill('solid', fgColor='FFFF00')
THIN = Side(style='thin', color='D0D3D8')
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
KEY_INPUTS = {'p_robot', 'p_rr', 'p_rent', 'p_care', 'bom', 'attach', 'rp', 'rd', 'rt', 'smr', 'option_rate', 'visit_cost', 'comm_cost'}
SCOL = {'C': 'D', 'B': 'E', 'U': 'F'}          # Inputs sheet scenario columns
YCOLS = ['E', 'F', 'G', 'H', 'I']               # Inputs_Yearly Y1..Y5
FCOLS = ['C', 'D', 'E', 'F', 'G']               # FM sheets Y1..Y5

wb = Workbook()

def setw(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[CL(i + 1)].width = w

def hdr(ws, row, vals):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=i + 1, value=v); c.font = HDR; c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def fmt_for(unit):
    return PCT if unit in ('%', '%/년') else (NUM1 if unit in ('만원/월', '만원/Kit', '만원/h', 'x', '회/년', '회/대·년') else NUM)

# ---------------------------------------------------------------- README
ws = wb.active; ws.title = 'README'
setw(ws, [3, 110])
lines = [
    ('MH Robotics — Seed + TIPS IR Financial Model (Draft v5, 2026-10)', TITLE),
    ('', None),
    ('목적: Seed · TIPS IR 덱의 모든 사업 수치를 하나의 수식 모델로 관리. 단위는 별도 표기가 없으면 만원 (KRW 10,000). 억원 = 만원 / 10,000.', None),
    ('주의: 실적 · 계약 · 고객 · Partner 데이터 없음. 모든 수치는 FACT / DERIVED / ASSUMPTION / TARGET Tag로 구분.', BOLD),
    ('', None),
    ('Color code', BOLD),
    ('  파란 글자 = 입력값 (Inputs, Inputs_Yearly, Budget_24M 팀 · 지출 표에서만 수정) · 검은 글자 = 수식 · 녹색 글자 = 다른 Sheet 참조 · 노란 바탕 = 핵심 가정 (WTP · BOM · Volume)', None),
    ('Tag', BOLD),
    ('  FACT = 공식 통계 · 공개자료 확인값 · DERIVED = FACT 또는 가정으로 계산한 값 · ASSUMPTION = 현재 사업가설 · TARGET = 24개월 · 이후 목표', None),
    ('Sheets', BOLD),
    ('  Inputs: 단일값 가정 (시나리오별 Conservative / Base / Upside 열)', None),
    ('  Inputs_Yearly: 연도별 가정 (Volume, BOM, Interface 표준부품 사용률, 인원 · 지출 등). Y1~Y2 지출 = Budget_24M 팀 · 지출 계획과 같은 값', None),
    ('  FM_Conservative / FM_Base / FM_Upside: 5개년 Bottom-up 손익 (INSTALL · OPERATE · EXPAND 매출층, 채널 3종)', None),
    ('  Scenario_Summary: 3개 시나리오 요약 (억원)', None),
    ('  Household: 대표 세대 5년 경제성 (Remodeling 구매 · Rental, Retrofit, New-build)', None),
    ('  Unit_Economics: Rental 월 원가 Build-up · Payback, Care · Consumables 단위 경제성, BOM 구성', None),
    ('  Market: Bottom-up 4개 시장 (Remodeling · Retrofit · New-build · Recurring)', None),
    ('  Sensitivity: 세대 Contribution 민감도 (수식) + 회사 Y5 Contribution 민감도 (model.py 정적 결과)', None),
    ('  Budget_24M: 24개월 팀 계획 · 사용처 · TIPS 과제 편성 · 재원 · Seed 범위', None),
    ('  Sources: FACT 출처', None),
    ('', None),
    ('Year 정의: Y1 = M1~M12 (Seed + TIPS 1차년도), Y2 = M13~M24, Y3 = Series A 이후 첫 해. 신축 Option은 계약 2년 후 입주 · 설치로 인식.', None),
    ('재현: python3 MH/source/model.py → python3 MH/source/xlsx_model.py (LibreOffice 재계산 후 check_xlsx.py로 model.py와 교차검증)', None),
]
for i, (t, f) in enumerate(lines):
    c = ws.cell(row=i + 1, column=2, value=t); c.font = f or BLACK
    c.alignment = Alignment(wrap_text=True, vertical='top')

# ---------------------------------------------------------------- Inputs (scalar) / Inputs_Yearly
wi = wb.create_sheet('Inputs'); wy = wb.create_sheet('Inputs_Yearly')
setw(wi, [16, 46, 12, 13, 13, 13, 12, 70]); setw(wy, [16, 46, 12, 13, 11, 11, 11, 11, 11, 12, 60])
wi['A1'] = 'Inputs — 단일값 가정 (만원 기준)'; wi['A1'].font = TITLE
wy['A1'] = 'Inputs_Yearly — 연도별 가정'; wy['A1'].font = TITLE
hdr(wi, 3, ['Key', '항목', '단위', 'Conservative', 'Base', 'Upside', 'Tag', 'Source / Note'])
hdr(wy, 3, ['Key', '항목', '단위', 'Scenario', 'Y1', 'Y2', 'Y3', 'Y4', 'Y5', 'Tag', 'Source / Note'])
REF = {}        # (key, s) -> absolute ref (scalar)  |  YREF[(key, s)] -> row in Inputs_Yearly
YREF = {}
ri, ry = 4, 4
grp_prev = None
for d in MD.INPUTS:
    if not d['yearly']:
        if d['group'] != grp_prev:
            c = wi.cell(row=ri, column=1, value=d['group'].upper()); c.font = BOLD
            for col in range(1, 9): wi.cell(row=ri, column=col).fill = SEC_FILL
            ri += 1; grp_prev = d['group']
        wi.cell(row=ri, column=1, value=d['key']).font = BLACK
        wi.cell(row=ri, column=2, value=d['desc']).font = BLACK
        wi.cell(row=ri, column=3, value=d['unit']).font = BLACK
        for s in MD.SC:
            c = wi[f'{SCOL[s]}{ri}']; c.value = d['vals'][s]; c.font = BLUE; c.number_format = fmt_for(d['unit'])
            if d['key'] in KEY_INPUTS: c.fill = KEY_FILL
            REF[(d['key'], s)] = f"Inputs!${SCOL[s]}${ri}"
        wi.cell(row=ri, column=7, value=d['tag']).font = BOLD
        wi.cell(row=ri, column=8, value=d['src']).font = BLACK
        ri += 1
    else:
        for s in MD.SC:
            wy.cell(row=ry, column=1, value=d['key']).font = BLACK
            wy.cell(row=ry, column=2, value=d['desc']).font = BLACK
            wy.cell(row=ry, column=3, value=d['unit']).font = BLACK
            wy.cell(row=ry, column=4, value=MD.SCN[s]).font = BLACK
            for t in range(5):
                c = wy[f'{YCOLS[t]}{ry}']; c.value = d['vals'][s][t]; c.font = BLUE; c.number_format = fmt_for(d['unit'])
                if d['key'] in KEY_INPUTS: c.fill = KEY_FILL
            wy.cell(row=ry, column=10, value=d['tag']).font = BOLD
            wy.cell(row=ry, column=11, value=d['src']).font = BLACK
            YREF[(d['key'], s)] = ry
            ry += 1
        ry += 1
wi.freeze_panes = 'C4'; wy.freeze_panes = 'E4'

def S(k, s):
    return REF[(k, s)]

def Y(k, s, t):
    return f"Inputs_Yearly!${YCOLS[t]}${YREF[(k, s)]}"

# ---------------------------------------------------------------- FM sheets
LINES = [
    # (name, label, unit, kind)  kind: 'h' header | 'n' number | 'p' percent
    ('H1', 'VOLUME', '', 'h'),
    ('rd', 'Remodeling — MH 직접 판매', '세대', 'n'),
    ('rp', 'Remodeling — Partner 경유', '세대', 'n'),
    ('rt', 'Existing Kitchen Retrofit', '세대', 'n'),
    ('ni', 'New-build Interface Option 입주 (계약 2년 후)', '세대', 'n'),
    ('kitchens', '설치 세대 합계', '세대', 'n'),
    ('att', '로봇 동시 구매율 (Remodeling)', '%', 'p'),
    ('later', 'Interface 선설치 세대 Robot 후설치', '대', 'n'),
    ('pl_remodel', 'Robot 설치 — Remodeling', '대', 'n'),
    ('pl_new', 'Robot 설치 — New-build 입주', '대', 'n'),
    ('pl', 'Robot 설치 합계 (Retrofit 포함)', '대', 'n'),
    ('rs', 'Rental 비중', '%', 'p'),
    ('pf', 'Rental Partner 보유 Flag', 'flag', 'n'),
    ('rpl', 'Rental 설치', '대', 'n'),
    ('pp', '구매 설치', '대', 'n'),
    ('drpl', 'Rental — MH 보유 (Pilot)', '대', 'n'),
    ('prpl', 'Rental — Partner 보유', '대', 'n'),
    ('pool', 'Interface 선설치 (Robot 미설치) 누적', '세대', 'n'),
    ('pb_end', '구매 Robot 누적 (기말)', '대', 'n'),
    ('pb_avg', '구매 Robot 평균 가동', '대', 'n'),
    ('dr_end', 'MH Rental 누적 (기말)', '대', 'n'),
    ('dr_avg', 'MH Rental 평균 가동', '대', 'n'),
    ('pr_end', 'Partner Rental 누적 (기말)', '대', 'n'),
    ('pr_avg', 'Partner Rental 평균 가동', '대', 'n'),
    ('base_end', 'Installed Robot (기말)', '대', 'n'),
    ('backlog', '신축 계약 Backlog (미설치 세대)', '세대', 'n'),
    ('H2', 'REVENUE (만원) — INSTALL / OPERATE / EXPAND', '', 'h'),
    ('real', '가격 실현율', '%', 'p'),
    ('cons_y', 'Consumables 연 List (구매 고객)', '만원/대', 'n'),
    ('cons_y_rent', 'Consumables 연 List (Rental, Grip 제외)', '만원/대', 'n'),
    ('grip_y', 'Grip Kit 연 List', '만원/대', 'n'),
    ('rev_kitchen', 'INSTALL. Interface · Integration', '만원', 'n'),
    ('rev_comm', 'INSTALL. Installation · Calibration', '만원', 'n'),
    ('rev_robot', 'INSTALL. Robot System (구매 + Rental Partner 공급)', '만원', 'n'),
    ('rev_rental', 'OPERATE. Rental (MH 보유)', '만원', 'n'),
    ('rev_care', 'OPERATE. Care (구매 고객 + Partner 서비스료)', '만원', 'n'),
    ('rev_cons', 'OPERATE. Consumables', '만원', 'n'),
    ('rev_upg', 'EXPAND. Skill + Tool', '만원', 'n'),
    ('rev', '매출 합계', '만원', 'n'),
    ('H3', 'COGS (만원)', '', 'h'),
    ('smr', 'Interface 표준부품 사용률', '%', 'p'),
    ('kit_unit', 'Interface Kit 원가/세대 (Remodeling)', '만원', 'n'),
    ('bom', 'Robot BOM/대', '만원', 'n'),
    ('care_unit', 'Care 원가/대·년', '만원', 'n'),
    ('c_kitchen', 'Interface 원가 (3개 채널)', '만원', 'n'),
    ('c_log', '물류', '만원', 'n'),
    ('c_robot', 'Robot BOM (판매분)', '만원', 'n'),
    ('c_comm', 'Installation · Calibration 원가', '만원', 'n'),
    ('c_warranty', 'Warranty Reserve', '만원', 'n'),
    ('capex', 'Rental 자산 취득 (MH 보유)', '만원', 'n'),
    ('cum_capex', 'Rental 자산 누적 취득', '만원', 'n'),
    ('c_dep', 'Rental 자산 감가상각', '만원', 'n'),
    ('c_care', 'Care Service 원가', '만원', 'n'),
    ('c_cons', 'Consumables 원가 (Rental Grip 포함)', '만원', 'n'),
    ('c_upg', 'Skill · Tool 원가', '만원', 'n'),
    ('cogs', 'COGS 합계', '만원', 'n'),
    ('gp', '매출총이익', '만원', 'n'),
    ('gm', '매출총이익률', '%', 'p'),
    ('H4', 'CHANNEL · OPEX (만원)', '', 'h'),
    ('ch_partner', 'Partner 수수료 (Remodeling Partner · Retrofit)', '만원', 'n'),
    ('ch_cac', '직접판매 획득비용', '만원', 'n'),
    ('ch_bd', '신축 Project 수주비용', '만원', 'n'),
    ('contrib', 'Contribution (매출총이익 − 채널 · 변동 판매비)', '만원', 'n'),
    ('op_people', '인건비', '만원', 'n'),
    ('op_other', '기타 운영비 (Robot HW · Hand · Mock-up · SW · Pilot · 고객 · 인증 · IP · 공간 · 관리 · 예비비 · 연구수당)', '만원', 'n'),
    ('opex', 'Opex 합계', '만원', 'n'),
    ('op', '영업이익 (근사)', '만원', 'n'),
    ('cash', '현금흐름 (영업이익 + 감가 − Rental 자산)', '만원', 'n'),
    ('cum_cash', '누적 현금흐름', '만원', 'n'),
    ('H5', 'MIX', '', 'h'),
    ('install', 'INSTALL 매출 (Interface + 설치 + Robot)', '만원', 'n'),
    ('build', 'Interface + 설치 매출', '만원', 'n'),
    ('recurring', 'OPERATE 매출 (Rental + Care + Consumables)', '만원', 'n'),
    ('expand', 'EXPAND 매출 (Skill + Tool)', '만원', 'n'),
    ('oe_share', 'OPERATE + EXPAND 비중', '%', 'p'),
    ('rr_share', 'Robot + Recurring + Upgrade 비중', '%', 'p'),
]
FM_ROW = {}

def fm_formulas(s, name, t, R, P, P2):
    """Return formula string for line `name`, year t (0-based). R(n) = this year's cell, P(n) = previous year or 0."""
    yv = lambda k: Y(k, s, t)
    sv = lambda k: S(k, s)
    f = {
        'rd': lambda: f"={yv('rd')}",
        'rp': lambda: f"={yv('rp')}",
        'ni': lambda: (f"={Y('projects', s, t - 2)}*{sv('hh_project')}*{sv('option_rate')}" if t >= 2 else "=0"),
        'rt': lambda: f"={yv('rt')}",
        'kitchens': lambda: f"={R('rd')}+{R('rp')}+{R('ni')}+{R('rt')}",
        'att': lambda: f"={yv('attach')}",
        'later': lambda: f"={P('pool')}*{sv('later_attach')}",
        'pl_remodel': lambda: f"=({R('rd')}+{R('rp')})*{R('att')}",
        'pl_new': lambda: f"={R('ni')}*{sv('new_attach')}",
        'pl': lambda: f"={R('pl_remodel')}+{R('pl_new')}+{R('later')}+{R('rt')}",
        'rs': lambda: f"={yv('rental_share')}",
        'pf': lambda: f"={yv('partner_rental')}",
        'rpl': lambda: f"={R('pl')}*{R('rs')}",
        'pp': lambda: f"={R('pl')}-{R('rpl')}",
        'drpl': lambda: f"={R('rpl')}*(1-{R('pf')})",
        'prpl': lambda: f"={R('rpl')}*{R('pf')}",
        'pool': lambda: f"={P('pool')}+({R('rd')}+{R('rp')})*(1-{R('att')})+{R('ni')}*(1-{sv('new_attach')})-{R('later')}",
        'pb_end': lambda: f"={P('pb_end')}+{R('pp')}",
        'pb_avg': lambda: f"={P('pb_end')}+{R('pp')}/2",
        'dr_end': lambda: f"={P('dr_end')}+{R('drpl')}",
        'dr_avg': lambda: f"={P('dr_end')}+{R('drpl')}/2",
        'pr_end': lambda: f"={P('pr_end')}+{R('prpl')}",
        'pr_avg': lambda: f"={P('pr_end')}+{R('prpl')}/2",
        'base_end': lambda: f"={R('pb_end')}+{R('dr_end')}+{R('pr_end')}",
        'backlog': lambda: (f"=({yv('projects')}+{Y('projects', s, t - 1)})*{sv('hh_project')}*{sv('option_rate')}" if t >= 1
                            else f"={yv('projects')}*{sv('hh_project')}*{sv('option_rate')}"),
        'real': lambda: f"={yv('realization')}",
        'cons_y': lambda: f"={sv('p_grip')}*{sv('n_grip')}+{sv('p_clean')}*{sv('n_clean')}+{sv('p_protect')}*{sv('n_protect')}",
        'cons_y_rent': lambda: f"={sv('p_clean')}*{sv('n_clean')}+{sv('p_protect')}*{sv('n_protect')}",
        'grip_y': lambda: f"={sv('p_grip')}*{sv('n_grip')}",
        'rev_kitchen': lambda: f"=({R('rd')}+{R('rp')})*{sv('p_rr')}*{R('real')}+{R('ni')}*{sv('p_rr_new')}+{R('rt')}*{sv('p_rt_if')}*{R('real')}",
        'rev_comm': lambda: f"=({R('pl')}-{R('rt')})*{sv('p_comm')}*{R('real')}+{R('rt')}*{sv('p_comm_rt')}*{R('real')}",
        'rev_robot': lambda: f"={R('pp')}*{sv('p_robot')}*{R('real')}+{R('prpl')}*{sv('p_robot')}*{sv('wholesale')}",
        'rev_rental': lambda: f"={R('dr_avg')}*{sv('p_rent')}*12",
        'rev_care': lambda: f"={R('pb_avg')}*{sv('care_attach')}*{sv('p_care')}+{R('pr_avg')}*{sv('partner_fee')}*12",
        'rev_cons': lambda: f"={R('pb_avg')}*{sv('cons_attach')}*{R('cons_y')}+({R('dr_avg')}+{R('pr_avg')})*{sv('cons_attach')}*{R('cons_y_rent')}",
        'rev_upg': lambda: f"={P('pl')}*{sv('sw_attach')}*{sv('p_sw')}+{P2('pl')}*{sv('tool_attach')}*{sv('p_tool')}",
        'rev': lambda: f"=SUM({R('rev_kitchen')}:{R('rev_upg')})",
        'smr': lambda: f"={yv('smr')}",
        'kit_unit': lambda: f"={sv('kit_std_cost')}*({R('smr')}+{sv('custom_factor')}*(1-{R('smr')}))+{sv('design_cost')}*(1-{R('smr')})",
        'bom': lambda: f"={yv('bom')}",
        'care_unit': lambda: f"={yv('visits')}*{yv('visit_cost')}+{sv('corrective')}*{sv('corr_cost')}+{sv('cloud')}",
        'c_kitchen': lambda: f"=({R('rd')}+{R('rp')})*{R('kit_unit')}+{R('ni')}*{sv('kit_new_cost')}+{R('rt')}*{sv('rt_kit_cost')}",
        'c_log': lambda: f"=({R('rd')}+{R('rp')}+{R('ni')}+{R('rt')})*{sv('logistics')}",
        'c_robot': lambda: f"=({R('pp')}+{R('prpl')})*{R('bom')}",
        'c_comm': lambda: f"=({R('pl')}-{R('rt')})*{yv('comm_cost')}+{R('rt')}*{yv('comm_cost_rt')}",
        'c_warranty': lambda: f"={sv('warranty')}*{R('rev_robot')}",
        'capex': lambda: f"={R('drpl')}*{R('bom')}",
        'cum_capex': lambda: f"={P('cum_capex')}+{R('capex')}",
        'c_dep': lambda: f"={P('cum_capex')}*(1-{sv('residual')})/5+{R('capex')}*(1-{sv('residual')})/5*0.5",
        'c_care': lambda: f"=({R('pb_avg')}*{sv('care_attach')}+{R('dr_avg')}+{R('pr_avg')})*{R('care_unit')}",
        'c_cons': lambda: f"={R('rev_cons')}*{sv('cons_cogs')}+({R('dr_avg')}+{R('pr_avg')})*{R('grip_y')}*{sv('cons_cogs')}",
        'c_upg': lambda: f"={P('pl')}*{sv('sw_attach')}*{sv('p_sw')}*{sv('sw_cogs')}+{P2('pl')}*{sv('tool_attach')}*{sv('p_tool')}*{sv('tool_cogs')}",
        'cogs': lambda: f"={R('c_kitchen')}+{R('c_log')}+{R('c_robot')}+{R('c_comm')}+{R('c_warranty')}+{R('c_dep')}+{R('c_care')}+{R('c_cons')}+{R('c_upg')}",
        'gp': lambda: f"={R('rev')}-{R('cogs')}",
        'gm': lambda: f"=IF({R('rev')}=0,0,{R('gp')}/{R('rev')})",
        'ch_partner': lambda: f"={sv('partner_margin')}*({R('rp')}*({sv('p_rr')}*{R('real')}+{R('att')}*({sv('p_robot')}+{sv('p_comm')})*{R('real')})+{R('rt')}*({sv('p_rt_if')}+{sv('p_robot')}+{sv('p_comm_rt')})*{R('real')})",
        'ch_cac': lambda: f"={R('rd')}*{sv('cac')}",
        'ch_bd': lambda: f"={yv('projects')}*{sv('bd_new')}",
        'contrib': lambda: f"={R('gp')}-{R('ch_partner')}-{R('ch_cac')}-{R('ch_bd')}",
        'op_people': lambda: f"={yv('fte')}*{yv('loaded')}",
        'op_other': lambda: '=' + '+'.join(yv(k) for k in MD.OTHER_KEYS),
        'opex': lambda: f"={R('op_people')}+{R('op_other')}",
        'op': lambda: f"={R('contrib')}-{R('opex')}",
        'cash': lambda: f"={R('op')}+{R('c_dep')}-{R('capex')}",
        'cum_cash': lambda: f"={P('cum_cash')}+{R('cash')}",
        'build': lambda: f"={R('rev_kitchen')}+{R('rev_comm')}",
        'install': lambda: f"={R('rev_kitchen')}+{R('rev_comm')}+{R('rev_robot')}",
        'recurring': lambda: f"={R('rev_rental')}+{R('rev_care')}+{R('rev_cons')}",
        'expand': lambda: f"={R('rev_upg')}",
        'oe_share': lambda: f"=IF({R('rev')}=0,0,({R('recurring')}+{R('expand')})/{R('rev')})",
        'rr_share': lambda: f"=IF({R('rev')}=0,0,({R('rev_robot')}+{R('recurring')}+{R('rev_upg')})/{R('rev')})",
    }[name]
    return f()

def build_fm(s):
    ws = wb.create_sheet(f'FM_{MD.SCN[s]}')
    setw(ws, [52, 10, 14, 14, 14, 14, 14, 15])
    ws['A1'] = f'5-Year Bottom-up Model — {MD.SCN[s]} (만원)'; ws['A1'].font = TITLE
    ws['A2'] = f'모든 입력은 Inputs / Inputs_Yearly의 {MD.SCN[s]} 열·행을 참조 (녹색). 영업이익은 이자·세금 제외 근사.'
    ws['A2'].font = BLACK
    hdr(ws, 3, ['항목', '단위'] + MD.YEARS + ['5Y 합계'])
    row = 4; rows = {}
    for name, label, unit, kind in LINES:
        rows[name] = row; row += 1
    FM_ROW[s] = rows
    sumable = {'rd', 'rp', 'rt', 'ni', 'kitchens', 'later', 'pl_remodel', 'pl_new', 'pl', 'rpl', 'pp', 'drpl', 'prpl',
               'rev_kitchen', 'rev_comm', 'rev_robot', 'rev_rental', 'rev_care', 'rev_cons', 'rev_upg', 'rev',
               'c_kitchen', 'c_log', 'c_robot', 'c_comm', 'c_warranty', 'capex', 'c_dep', 'c_care', 'c_cons', 'c_upg',
               'cogs', 'gp', 'ch_partner', 'ch_cac', 'ch_bd', 'contrib', 'op_people', 'op_other', 'opex', 'op', 'cash',
               'build', 'recurring', 'install', 'expand'}
    for name, label, unit, kind in LINES:
        r = rows[name]
        if kind == 'h':
            c = ws.cell(row=r, column=1, value=label); c.font = BOLD
            for col in range(1, 9): ws.cell(row=r, column=col).fill = SEC_FILL
            continue
        ws.cell(row=r, column=1, value=label).font = BOLD if name in ('rev', 'gp', 'contrib', 'op', 'kitchens', 'pl', 'oe_share') else BLACK
        ws.cell(row=r, column=2, value=unit).font = BLACK
        for t in range(5):
            col = FCOLS[t]
            R = lambda n, col=col: f"{col}{rows[n]}"
            P = lambda n, t=t: (f"{FCOLS[t - 1]}{rows[n]}" if t >= 1 else "0")
            P2 = lambda n, t=t: (f"{FCOLS[t - 2]}{rows[n]}" if t >= 2 else "0")
            fstr = fm_formulas(s, name, t, R, P, P2)
            c = ws[f'{col}{r}']; c.value = fstr
            c.font = GREEN if ('Inputs' in fstr and fstr.count('!') >= 1 and all(op not in fstr.split('!')[0] for op in '+-*/(')) else BLACK
            c.number_format = PCT if kind == 'p' else NUM
        if name in sumable:
            c = ws[f'H{r}']; c.value = f"=SUM(C{r}:G{r})"; c.font = BLACK; c.number_format = NUM
    ws.freeze_panes = 'C4'
    return ws

for s in MD.SC:
    build_fm(s)

def fmref(s, name, t):
    return f"FM_{MD.SCN[s]}!${FCOLS[t]}${FM_ROW[s][name]}"

# ---------------------------------------------------------------- Scenario_Summary (억원)
ws = wb.create_sheet('Scenario_Summary', 1)
setw(ws, [44, 14, 12, 12, 12, 12, 12])
ws['A1'] = 'Scenario Summary (억원 · 세대 · 대)'; ws['A1'].font = TITLE
ws['A2'] = '각 셀은 FM Sheet 참조. 억원 = 만원 / 10,000. Upside는 가격 인상 없이 Partner·표준화·설치원가·Installed Base로 차이.'
r = 4
for s in MD.SC:
    hdr(ws, r, [MD.SCN[s], '단위'] + MD.YEARS); r += 1
    for name, label, unit, div in [('rev', '매출', '억원', 10000), ('gp', '매출총이익', '억원', 10000), ('gm', '매출총이익률', '%', None),
                                   ('contrib', 'Contribution', '억원', 10000), ('opex', 'Opex', '억원', 10000),
                                   ('op', '영업이익 (근사)', '억원', 10000), ('cum_cash', '누적 현금흐름', '억원', 10000),
                                   ('kitchens', '설치 세대 (Remodeling + Retrofit + New-build)', '세대', None), ('pl', 'Robot 설치', '대', None),
                                   ('base_end', 'Installed Robot (기말)', '대', None), ('backlog', '신축 Backlog', '세대', None),
                                   ('oe_share', 'OPERATE + EXPAND 비중', '%', None)]:
        ws.cell(row=r, column=1, value=label).font = BLACK
        ws.cell(row=r, column=2, value=unit).font = BLACK
        for t in range(5):
            c = ws.cell(row=r, column=3 + t, value=f"={fmref(s, name, t)}" + (f"/{div}" if div else ''))
            c.font = GREEN; c.number_format = PCT if unit == '%' else (NUM1 if unit == '억원' else NUM)
        r += 1
    r += 1

# ---------------------------------------------------------------- Household (Base)
ws = wb.create_sheet('Household')
setw(ws, [46, 15, 15, 15, 15, 15, 15, 44])
ws['A1'] = 'Household Economics — 대표 세대, 5년, Base (만원)'; ws['A1'].font = TITLE
ws['A2'] = ('기대값 기준 (Skill · Tool은 구매율 반영). 구매 모델은 Care 가입 세대 기준. Y3 / Y5 = 해당 연도의 원가 수준을 5년간 적용. '
            'Remodeling = 직접판매 (획득비용), Retrofit = Partner 경유 (수수료), New-build = Project 수주비용 ÷ Option 세대.')
b = lambda k: S(k, 'B')
yb = lambda k, t: Y(k, 'B', t)
# column spec: (col, header, mode, install, t)
HH_COLS = [('B', 'Remodeling 구매 · Y3', 'p', 'remodel', 2), ('C', 'Remodeling 구매 · Y5', 'p', 'remodel', 4),
           ('D', 'Remodeling Rental · Y3', 'r', 'remodel', 2), ('E', 'Remodeling Rental · Y5', 'r', 'remodel', 4),
           ('F', 'Retrofit 구매 · Y3', 'p', 'retrofit', 2), ('G', 'New-build 구매 · Y3', 'p', 'newbuild', 2)]
hdr(ws, 4, ['항목'] + [c[1] for c in HH_COLS] + ['산식'])
CONS_Y = f"({b('p_grip')}*{b('n_grip')}+{b('p_clean')}*{b('n_clean')}+{b('p_protect')}*{b('n_protect')})"
CONS_YR = f"({b('p_clean')}*{b('n_clean')}+{b('p_protect')}*{b('n_protect')})"
KIT = lambda t: f"{b('kit_std_cost')}*({yb('smr', t)}+{b('custom_factor')}*(1-{yb('smr', t)}))+{b('design_cost')}*(1-{yb('smr', t)})"
CARE = lambda t: f"({yb('visits', t)}*{yb('visit_cost', t)}+{b('corrective')}*{b('corr_cost')}+{b('cloud')})"
def hh_f(label, mode, inst, t, col, hrow):
    r_if = {'remodel': b('p_rr'), 'retrofit': b('p_rt_if'), 'newbuild': b('p_rr_new')}[inst]
    r_comm = b('p_comm_rt') if inst == 'retrofit' else b('p_comm')
    c_if = {'remodel': KIT(t), 'retrofit': b('rt_kit_cost'), 'newbuild': b('kit_new_cost')}[inst]
    c_comm = yb('comm_cost_rt', t) if inst == 'retrofit' else yb('comm_cost', t)
    R = lambda lab: f"{col}{hrow[lab]}"
    if label == 'INSTALL · Interface · Integration': return f"={r_if}"
    if label == 'INSTALL · Installation · Calibration': return f"={r_comm}"
    if label == 'INSTALL · Robot System (구매)': return f"={b('p_robot')}" if mode == 'p' else "=0"
    if label == 'OPERATE · Rental (60개월)': return f"={b('p_rent')}*{b('rent_months')}" if mode == 'r' else "=0"
    if label == 'OPERATE · Care Basic (5년)': return f"=5*{b('p_care')}" if mode == 'p' else "=0"
    if label == 'OPERATE · Consumables (5년, 구매율 반영)':
        return f"=5*{b('cons_attach')}*{CONS_Y}" if mode == 'p' else f"=5*{b('cons_attach')}*{CONS_YR}"
    if label == 'EXPAND · ASSIST Skill (기대값)': return f"={b('sw_attach')}*{b('p_sw')}"
    if label == 'EXPAND · Tool / End-effector (기대값)': return f"={b('tool_attach')}*{b('p_tool')}"
    if label == '5년 매출': return '=' + '+'.join(R(l) for l in REV_L)
    if label == 'Interface 원가': return f"={c_if}"
    if label == 'Installation · Calibration 원가': return f"={c_comm}"
    if label == '물류': return f"={b('logistics')}"
    if label == 'Robot BOM (Rental은 잔존가치 차감)': return f"={yb('bom', t)}" if mode == 'p' else f"={yb('bom', t)}*(1-{b('residual')})"
    if label == 'Rental 금융비용 (5년)': return "=0" if mode == 'p' else f"={yb('bom', t)}*(1+{b('residual')})/2*{b('fin_rate')}*5"
    if label == 'Warranty Reserve': return f"={b('warranty')}*{b('p_robot')}" if mode == 'p' else f"={b('warranty')}*{yb('bom', t)}"
    if label == 'Care Service 원가 (5년)': return f"=5*{CARE(t)}"
    if label == 'Consumables 원가':
        return (f"={R('OPERATE · Consumables (5년, 구매율 반영)')}*{b('cons_cogs')}" if mode == 'p'
                else f"=({R('OPERATE · Consumables (5년, 구매율 반영)')}+5*{b('p_grip')}*{b('n_grip')})*{b('cons_cogs')}")
    if label == 'Skill · Tool 원가': return f"={b('sw_attach')}*{b('p_sw')}*{b('sw_cogs')}+{b('tool_attach')}*{b('p_tool')}*{b('tool_cogs')}"
    if label == '채널비용 (획득비용 · Partner 수수료 · 수주비용)':
        if inst == 'newbuild': return f"={b('bd_new')}/({b('hh_project')}*{b('option_rate')})"
        if inst == 'retrofit': return f"={b('partner_margin')}*({r_if}+{b('p_robot')}+{r_comm})"
        return f"={b('cac')}"
    if label == '5년 비용': return '=' + '+'.join(R(l) for l in COST_L)
    if label == 'Year 0 매출 (INSTALL)': return f"={R('INSTALL · Interface · Integration')}+{R('INSTALL · Installation · Calibration')}+{R('INSTALL · Robot System (구매)')}"
    if label == 'OPERATE 매출 (5년)': return f"={R('OPERATE · Rental (60개월)')}+{R('OPERATE · Care Basic (5년)')}+{R('OPERATE · Consumables (5년, 구매율 반영)')}"
    if label == 'EXPAND 매출 (5년)': return f"={R('EXPAND · ASSIST Skill (기대값)')}+{R('EXPAND · Tool / End-effector (기대값)')}"
    if label == '5년 매출총이익 (채널비용 전)': return f"={R('5년 매출')}-({R('5년 비용')}-{R('채널비용 (획득비용 · Partner 수수료 · 수주비용)')})"
    if label == '5년 Lifetime Contribution': return f"={R('5년 매출')}-{R('5년 비용')}"
    if label == 'Contribution Margin': return f"=IF({R('5년 매출')}=0,0,{R('5년 Lifetime Contribution')}/{R('5년 매출')})"
    if label == '5년 서비스 원가 (Care · Consumables · Warranty)': return f"={R('Care Service 원가 (5년)')}+{R('Consumables 원가')}+{R('Warranty Reserve')}"
    raise KeyError(label)
REV_L = ['INSTALL · Interface · Integration', 'INSTALL · Installation · Calibration', 'INSTALL · Robot System (구매)',
         'OPERATE · Rental (60개월)', 'OPERATE · Care Basic (5년)', 'OPERATE · Consumables (5년, 구매율 반영)',
         'EXPAND · ASSIST Skill (기대값)', 'EXPAND · Tool / End-effector (기대값)']
COST_L = ['Interface 원가', 'Installation · Calibration 원가', '물류', 'Robot BOM (Rental은 잔존가치 차감)', 'Rental 금융비용 (5년)',
          'Warranty Reserve', 'Care Service 원가 (5년)', 'Consumables 원가', 'Skill · Tool 원가', '채널비용 (획득비용 · Partner 수수료 · 수주비용)']
HH = [('REVENUE', None)] + [(l, '') for l in REV_L] + [('5년 매출', '')] + [('COST', None)] + [(l, '') for l in COST_L] + \
     [('5년 비용', ''), ('RESULT', None), ('Year 0 매출 (INSTALL)', ''), ('OPERATE 매출 (5년)', ''), ('EXPAND 매출 (5년)', ''),
      ('5년 매출총이익 (채널비용 전)', ''), ('5년 Lifetime Contribution', ''), ('Contribution Margin', ''),
      ('5년 서비스 원가 (Care · Consumables · Warranty)', '')]
r0 = 5; hrow = {}
for i, (label, _) in enumerate(HH):
    hrow[label] = r0 + i
for label, kind in HH:
    r = hrow[label]
    if kind is None:
        c = ws.cell(row=r, column=1, value=label); c.font = BOLD
        for col in range(1, 9): ws.cell(row=r, column=col).fill = SEC_FILL
        continue
    ws.cell(row=r, column=1, value=label).font = BOLD if label.startswith('5년') or label == 'Contribution Margin' else BLACK
    for col, _, mode, inst, t in HH_COLS:
        fs = hh_f(label, mode, inst, t, col, hrow)
        c = ws[f'{col}{r}']; c.value = fs; c.number_format = PCT if label == 'Contribution Margin' else NUM1
        c.font = GREEN if fs.count('!') == 1 and fs.startswith('=Inputs') and all(o not in fs[1:] for o in '+-*/(') else BLACK

# ---------------------------------------------------------------- Unit_Economics (Base)
ws = wb.create_sheet('Unit_Economics')
setw(ws, [52, 16, 16, 50])
ws['A1'] = 'Unit Economics — Base (만원)'; ws['A1'].font = TITLE
hdr(ws, 3, ['Rental 월 원가 Build-up (Robot 1대, MH 보유 가정)', 'Y3 원가', 'Y5 원가', '산식'])
UE = [
    ('Robot BOM', lambda t: f"={yb('bom', t)}", ''),
    ('감가 (잔존가치 차감, 60개월)', lambda t: f"=B_bom*(1-{b('residual')})/{b('rent_months')}", ''),
    ('금융비용 (평균잔액 × 금리 / 12)', lambda t: f"=B_bom*(1+{b('residual')})/2*{b('fin_rate')}/12", ''),
    ('Care 원가 / 12', lambda t: f"=({yb('visits', t)}*{yb('visit_cost', t)}+{b('corrective')}*{b('corr_cost')}+{b('cloud')})/12", ''),
    ('Grip Kit 원가 / 12', lambda t: f"={b('p_grip')}*{b('n_grip')}*{b('cons_cogs')}/12", ''),
    ('Failure Reserve', lambda t: f"={b('warranty')}*B_bom/{b('rent_months')}", ''),
    ('월 원가 합계', 'SUM', ''),
    ('월 Rental 요금 (가정)', lambda t: f"={b('p_rent')}", ''),
    ('월 Contribution', 'CONTRIB', ''),
    ('Contribution Margin', 'MARGIN', ''),
    ('마진 20% 확보 요금', 'FEE20', ''),
    ('Payback (BOM ÷ (요금 − Care − Grip), 개월)', 'PAYBACK', '금융비용 제외 현금 회수'),
    ('Payback 허들 충족 최대 BOM', 'BOMMAX', 'Rental Partner 요구 Payback 기준'),
]
ue_row = {lab: 4 + i for i, (lab, _, _) in enumerate(UE)}
for lab, fn, note in UE:
    r = ue_row[lab]
    ws.cell(row=r, column=1, value=lab).font = BOLD if lab in ('월 원가 합계', '월 Contribution', 'Payback (BOM ÷ (요금 − Care − Grip), 개월)') else BLACK
    ws.cell(row=r, column=4, value=note).font = BLACK
    for col, t in (('B', 2), ('C', 4)):
        bomc = f"{col}{ue_row['Robot BOM']}"
        cost_rng = [ue_row[k] for k in ('감가 (잔존가치 차감, 60개월)', '금융비용 (평균잔액 × 금리 / 12)', 'Care 원가 / 12', 'Grip Kit 원가 / 12', 'Failure Reserve')]
        if fn == 'SUM': fs = '=' + '+'.join(f'{col}{x}' for x in cost_rng)
        elif fn == 'CONTRIB': fs = f"={col}{ue_row['월 Rental 요금 (가정)']}-{col}{ue_row['월 원가 합계']}"
        elif fn == 'MARGIN': fs = f"={col}{ue_row['월 Contribution']}/{col}{ue_row['월 Rental 요금 (가정)']}"
        elif fn == 'FEE20': fs = f"={col}{ue_row['월 원가 합계']}/0.8"
        elif fn == 'PAYBACK': fs = f"={bomc}/({col}{ue_row['월 Rental 요금 (가정)']}-{col}{ue_row['Care 원가 / 12']}-{col}{ue_row['Grip Kit 원가 / 12']})"
        elif fn == 'BOMMAX': fs = f"={b('payback_hurdle')}*({col}{ue_row['월 Rental 요금 (가정)']}-{col}{ue_row['Care 원가 / 12']}-{col}{ue_row['Grip Kit 원가 / 12']})"
        else: fs = fn(t).replace('B_bom', bomc)
        c = ws[f'{col}{r}']; c.value = fs; c.font = BLACK
        c.number_format = PCT if lab == 'Contribution Margin' else NUM1
r = ue_row['Payback 허들 충족 최대 BOM'] + 2
hdr(ws, r, ['Care Basic (Robot 1대·년)', 'Y3 원가', 'Y5 원가', '산식']); r += 1
care_r = r
for lab, fn in [('Care 연 요금', lambda t: f"={b('p_care')}"),
                ('정기 방문 원가', lambda t: f"={yb('visits', t)}*{yb('visit_cost', t)}"),
                ('고장 방문 원가', lambda t: f"={b('corrective')}*{b('corr_cost')}"),
                ('Cloud·Software', lambda t: f"={b('cloud')}"),
                ('Care 원가 합계', None), ('Care Contribution', None), ('Care Margin', None)]:
    ws.cell(row=r, column=1, value=lab).font = BOLD if fn is None else BLACK
    for col, t in (('B', 2), ('C', 4)):
        if lab == 'Care 원가 합계': fs = f"=SUM({col}{care_r + 1}:{col}{care_r + 3})"
        elif lab == 'Care Contribution': fs = f"={col}{care_r}-{col}{care_r + 4}"
        elif lab == 'Care Margin': fs = f"={col}{care_r + 5}/{col}{care_r}"
        else: fs = fn(t)
        c = ws[f'{col}{r}']; c.value = fs; c.font = BLACK; c.number_format = PCT if lab == 'Care Margin' else NUM1
    r += 1
r += 1
hdr(ws, r, ['Consumables (Robot 1대·년, 구매 고객)', '연 매출', '원가', 'Contribution']); r += 1
cons_r0 = r
for lab, p, n in [('Grip Kit', 'p_grip', 'n_grip'), ('Cleaning Kit', 'p_clean', 'n_clean'), ('Protection Kit', 'p_protect', 'n_protect')]:
    ws.cell(row=r, column=1, value=f'{lab} (List × 구매율)').font = BLACK
    ws[f'B{r}'] = f"={b(p)}*{b(n)}*{b('cons_attach')}"; ws[f'C{r}'] = f"=B{r}*{b('cons_cogs')}"; ws[f'D{r}'] = f"=B{r}-C{r}"
    for col in 'BCD': ws[f'{col}{r}'].number_format = NUM1; ws[f'{col}{r}'].font = BLACK
    r += 1
ws.cell(row=r, column=1, value='합계').font = BOLD
for col in 'BCD':
    ws[f'{col}{r}'] = f"=SUM({col}{cons_r0}:{col}{r - 1})"; ws[f'{col}{r}'].number_format = NUM1; ws[f'{col}{r}'].font = BOLD

r += 2
hdr(ws, r, ['Rental Partner 경제성 (Base, Y4~)', '값', '', '산식']); r += 1
pr0 = r
for lab, fs, nf, note in [
        ('Robot 매입가 (ASP × 공급가율)', f"={b('p_robot')}*{b('wholesale')}", NUM1, ''),
        ('월 순유입 (요금 − MH 서비스료)', f"={b('p_rent')}-{b('partner_fee')}", NUM1, ''),
        ('잔존가치 (매입가 × 잔존율)', f"=B{pr0}*{b('residual')}", NUM1, ''),
        ('월 IRR', f"=RATE({b('rent_months')},B{pr0 + 1},-B{pr0},B{pr0 + 2})", '0.00%', 'RATE(기간, 월 유입, −매입가, 잔존가치)'),
        ('연 IRR (연체·해지 미반영)', f"=(1+B{pr0 + 3})^12-1", PCT, ''),
        ('단순 회수기간 (개월)', f"=B{pr0}/B{pr0 + 1}", NUM1, '매입가 ÷ 월 순유입'),
        ('Partner 요구 Payback (가정, 개월)', f"={b('payback_hurdle')}", NUM1, '회수기간 > 요구치이면 매입가율 · 서비스료 · 기간 협의 필요')]:
    ws.cell(row=r, column=1, value=lab).font = BOLD if '연 IRR' in lab else BLACK
    c = ws.cell(row=r, column=2, value=fs); c.font = BLACK; c.number_format = nf
    ws.cell(row=r, column=4, value=note).font = BLACK
    r += 1
r += 1
hdr(ws, r, ['Robot BOM 구성 (ASSUMPTION, 만원/대)', 'Pilot (Y1)', 'Y3', 'Y5']); r += 1
bb0 = r
for item, p1, p3, p5, basis in MD.BOM_BREAKDOWN:
    ws.cell(row=r, column=1, value=item).font = BLACK
    for col, val in (('B', p1), ('C', p3), ('D', p5)):
        c = ws[f'{col}{r}']; c.value = val; c.font = BLUE; c.number_format = NUM
    ws.cell(row=r, column=5, value=basis).font = BLACK
    r += 1
ws.cell(row=r, column=1, value='합계').font = BOLD
for col in 'BCD':
    ws[f'{col}{r}'] = f"=SUM({col}{bb0}:{col}{r - 1})"; ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM
r += 1
ws.cell(row=r, column=1, value='검증: Inputs_Yearly Base BOM과 차이 (0이어야 함)').font = BLACK
for col, t in (('B', 0), ('C', 2), ('D', 4)):
    ws[f'{col}{r}'] = f"={col}{r - 1}-{Y('bom', 'B', t)}"; ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
ws.column_dimensions['E'].width = 60
UE_EXTRA = {'irr_row': pr0 + 4, 'bom_check_row': r}

# ---------------------------------------------------------------- Market
ws = wb.create_sheet('Market')
setw(ws, [60, 16, 12, 12, 64])
ws['A1'] = 'Market Sizing — Bottom-up 4개 시장 (Base)'; ws['A1'].font = TITLE
ws['A2'] = '큰 TAM 대신 세대 수 × 적용률 × 단가. 아파트 재고는 잠재 대상이지 실구매 시장 규모가 아님. 핵심 비율은 모두 ASSUMPTION (검증 계획 부록).'
hdr(ws, 3, ['항목', '값', '단위', 'Tag', '산식 / 출처'])
cons_y_f = f"({b('p_grip')}*{b('n_grip')}+{b('p_clean')}*{b('n_clean')}+{b('p_protect')}*{b('n_protect')})"
MK = [
    ('총주택 (2025)', f"={b('m_housing_total')}", '천호', 'FACT', '국가데이터처 2025 인구주택총조사'),
    ('아파트 비중', f"={b('m_apt_share')}", '%', 'FACT', ''),
    ('아파트 수 (2025)', '=B4*B5', '천호', 'DERIVED', '총주택 × 비중'),
    ('교차검증 ①: 20년+ 아파트(2023) ÷ 교체주기', f"={b('m_apt_20y_2023')}/{b('a_replace_cycle')}", '천/년', 'DERIVED', ''),
    ('교차검증 ②: 매매거래 × 아파트 비중 × 교체율 + 비거래 교체', f"={b('m_txn_2025')}*{b('a_txn_apt_share')}*{b('a_txn_kitchen_rate')}+{b('a_aging_nontxn')}", '천/년', 'DERIVED', ''),
    ('연간 Kitchen 교체 세대 (설정)', f"={b('a_kitchen_replace')}", '천/년', 'ASSUMPTION', '①·② 교차검증 후 설정'),
    ('Premium 비중', f"={b('a_premium_share')}", '%', 'ASSUMPTION', ''),
    ('Premium Kitchen 교체 세대', '=B9*B10', '천/년', 'DERIVED', ''),
    ('Remodeling 적용 가능률', f"={b('a_fit_rate')}", '%', 'ASSUMPTION', ''),
    ('① Remodeling 대상 세대', '=B11*B12', '천/년', 'DERIVED', ''),
    ('Remodeling 패키지 기대 매출/세대', f"={b('p_rr')}+{Y('attach', 'B', 4)}*({b('p_robot')}+{b('p_comm')})", '만원', 'DERIVED', 'Interface + 로봇 구매율 × (Robot + 설치)'),
    ('① SAM Remodeling', '=B13*B14/10', '억원/년', 'DERIVED', '천세대 × 만원 ÷ 10 = 억원'),
    ('Premium 세대 비중 (재고)', f"={b('a_prem_stock')}", '%', 'ASSUMPTION', ''),
    ('Premium 세대 식기세척기 보유율', f"={b('a_dw_premium')}", '%', 'ASSUMPTION', 'TO BE VALIDATED'),
    ('Retrofit 호환률', f"={b('a_retro_fit')}", '%', 'ASSUMPTION', '받은 평면 5종 중 2종 (표본 작음)'),
    ('Retrofit 호환 세대 Pool (재고)', '=B6*B16*B17*B18', '천세대', 'DERIVED', '아파트 × Premium × 식세기 × 호환'),
    ('Retrofit 연간 전환율', f"={b('a_retro_conv')}", '%', 'ASSUMPTION', ''),
    ('② Retrofit 대상 세대', '=B19*B20', '천/년', 'DERIVED', ''),
    ('Retrofit 패키지 매출/세대', f"={b('p_rt_if')}+{b('p_robot')}+{b('p_comm_rt')}", '만원', 'DERIVED', 'Interface Kit + Robot + 설치'),
    ('② SAM Retrofit', '=B21*B22/10', '억원/년', 'DERIVED', ''),
    ('연간 신규 아파트 입주', f"={b('a_new_supply')}", '천/년', 'DERIVED', '2025 23.6만 · 2026E 18.3만'),
    ('Premium 단지 비중', f"={b('a_premium_project')}", '%', 'ASSUMPTION', ''),
    ('신축 Premium 세대', '=B24*B25', '천/년', 'DERIVED', ''),
    ('Interface Option 선택률', f"={b('option_rate')}", '%', 'ASSUMPTION', ''),
    ('③ New-build 대상 세대', '=B26*B27', '천/년', 'DERIVED', ''),
    ('New-build 패키지 기대 매출/세대', f"={b('p_rr_new')}+{b('new_attach')}*({b('p_robot')}+{b('p_comm')})", '만원', 'DERIVED', 'Option + 입주 로봇 구매율 × (Robot + 설치)'),
    ('③ SAM New-build', '=B28*B29/10', '억원/년', 'DERIVED', ''),
    ('구매 고객 ARPU (Care × 가입률 + Consumables × 구매율)', f"={b('care_attach')}*{b('p_care')}+{b('cons_attach')}*{cons_y_f}", '만원/년', 'DERIVED', ''),
    ('④ Installed Base 1,000대당 연 OPERATE 매출 (구매 고객)', '=B31*1000/10000', '억원/년', 'DERIVED', ''),
    ('Y5 Installed Robot (Base Plan)', f"={fmref('B', 'base_end', 4)}", '대', 'TARGET', 'FM_Base Y5'),
    ('④ Y5 OPERATE 매출 (Base Plan)', f"={fmref('B', 'recurring', 4)}/10000", '억원/년', 'TARGET', 'FM_Base Y5 (Rental Partner 보유분은 서비스료만)'),
    ('전체 패키지 단가 (Interface + Robot + 설치)', f"={b('p_rr')}+{b('p_robot')}+{b('p_comm')}", '만원', 'ASSUMPTION', ''),
    ('참고 TAM: Premium 세대 (Remodeling + New-build) × 전체 패키지', '=(B11+B26)*B35/10', '억원/년', 'DERIVED', '본문 미사용 (참고)'),
    ('SAM 합계 (① + ② + ③)', '=B15+B23+B30', '억원/년', 'DERIVED', ''),
    ('SOM: Y5 매출 (Base Plan)', f"={fmref('B', 'rev', 4)}/10000", '억원/년', 'TARGET', 'FM_Base Y5'),
    ('SOM 세대 / SAM 세대', f"={fmref('B', 'kitchens', 4)}/((B13+B21+B28)*1000)", '%', 'DERIVED', ''),
]
MK_ROW = {}
for i, (lab, fs, unit, tag, note) in enumerate(MK):
    r = 4 + i; MK_ROW[lab] = r
    ws.cell(row=r, column=1, value=lab).font = BOLD if lab.startswith(('①', '②', '③', '④', 'SAM 합계', 'SOM')) else BLACK
    c = ws.cell(row=r, column=2, value=fs); c.number_format = PCT if unit == '%' else NUM1
    c.font = GREEN if fs.startswith('=Inputs') and fs.count('!') == 1 and all(o not in fs[1:] for o in '+-*/(') else BLACK
    ws.cell(row=r, column=3, value=unit).font = BLACK
    ws.cell(row=r, column=4, value=tag).font = BOLD
    ws.cell(row=r, column=5, value=note).font = BLACK
assert MK_ROW['아파트 수 (2025)'] == 6 and MK_ROW['① SAM Remodeling'] == 15 and MK_ROW['SOM 세대 / SAM 세대'] == 39

# ---------------------------------------------------------------- Sensitivity
ws = wb.create_sheet('Sensitivity')
PARAMS = ['p_robot', 'bom', 'p_rr', 'cac', 'p_care', 'visit_cost', 'smr', 'comm_cost', 'corrective', 'warranty', 'cons_attach']
setw(ws, [40] + [11] * len(PARAMS) + [14, 14])
ws['A1'] = 'Sensitivity — Remodeling 구매 세대당 5년 Contribution (Base, Y3 원가, 직접판매)'; ws['A1'].font = TITLE
ws['A2'] = '각 행은 한 변수만 바꾼 파라미터 세트 (파란 글자 = 변경값). Contribution 수식은 Household 시트와 동일 구조.'
hdr(ws, 4, ['Case'] + PARAMS + ['Contribution', 'Δ vs Base'])
base_vals = {'p_robot': f"={b('p_robot')}", 'bom': f"={yb('bom', 2)}", 'p_rr': f"={b('p_rr')}", 'cac': f"={b('cac')}",
             'p_care': f"={b('p_care')}", 'visit_cost': f"={yb('visit_cost', 2)}", 'smr': f"={yb('smr', 2)}",
             'comm_cost': f"={yb('comm_cost', 2)}", 'corrective': f"={b('corrective')}", 'warranty': f"={b('warranty')}",
             'cons_attach': f"={b('cons_attach')}"}
CASES = [('Base', {})]
for name, lo, hi in [('Robot ASP (WTP)', {'p_robot': '*0.8'}, {'p_robot': '*1.2'}), ('Robot BOM', {'bom': '*1.2'}, {'bom': '*0.8'}),
                     ('Interface 가격', {'p_rr': '*0.8'}, {'p_rr': '*1.2'}), ('직접판매 획득비용', {'cac': '*1.5'}, {'cac': '*0.5'}),
                     ('고장 횟수', {'corrective': '=1.2', 'warranty': '=0.06'}, {'corrective': '=0.3', 'warranty': '=0.025'}),
                     ('Care 요금', {'p_care': '*0.8'}, {'p_care': '*1.2'}), ('Care 방문 원가', {'visit_cost': '*1.3'}, {'visit_cost': '*0.7'}),
                     ('Interface 표준부품 사용률', {'smr': '-0.15'}, {'smr': '+0.15'}), ('설치 · Calibration 원가', {'comm_cost': '*1.5'}, {'comm_cost': '*0.5'}),
                     ('Consumables 구매율', {'cons_attach': '=0.5'}, {'cons_attach': '=0.9'})]:
    CASES.append((name + ' — 불리', lo)); CASES.append((name + ' — 유리', hi))
pc = {p: CL(2 + i) for i, p in enumerate(PARAMS)}
for i, (name, ch) in enumerate(CASES):
    r = 5 + i
    ws.cell(row=r, column=1, value=name).font = BOLD if name == 'Base' else BLACK
    for p in PARAMS:
        c = ws[f'{pc[p]}{r}']
        if p in ch:
            op = ch[p]
            c.value = (op if op.startswith('=') else f"{base_vals[p]}{op}"); c.font = BLUE
        else:
            c.value = base_vals[p]; c.font = GREEN
        c.number_format = PCT if p in ('smr', 'warranty', 'cons_attach') else NUM1
    g = lambda p: f'{pc[p]}{r}'
    cons_y = f"({b('p_grip')}*{b('n_grip')}+{b('p_clean')}*{b('n_clean')}+{b('p_protect')}*{b('n_protect')})"
    rev = (f"{g('p_rr')}+{b('p_comm')}+{g('p_robot')}+5*{g('p_care')}+5*{g('cons_attach')}*{cons_y}"
           f"+{b('sw_attach')}*{b('p_sw')}+{b('tool_attach')}*{b('p_tool')}")
    cost = (f"{b('kit_std_cost')}*({g('smr')}+{b('custom_factor')}*(1-{g('smr')}))+{b('design_cost')}*(1-{g('smr')})"
            f"+{g('comm_cost')}+{b('logistics')}+{g('bom')}+{g('warranty')}*{g('p_robot')}"
            f"+5*({yb('visits', 2)}*{g('visit_cost')}+{g('corrective')}*{b('corr_cost')}+{b('cloud')})"
            f"+5*{g('cons_attach')}*{cons_y}*{b('cons_cogs')}"
            f"+{b('sw_attach')}*{b('p_sw')}*{b('sw_cogs')}+{b('tool_attach')}*{b('p_tool')}*{b('tool_cogs')}+{g('cac')}")
    cc = CL(2 + len(PARAMS)); dc = CL(3 + len(PARAMS))
    ws[f'{cc}{r}'] = f"=({rev})-({cost})"; ws[f'{cc}{r}'].number_format = NUM1; ws[f'{cc}{r}'].font = BLACK
    ws[f'{dc}{r}'] = f"={cc}{r}-${cc}$5"; ws[f'{dc}{r}'].number_format = NUM1; ws[f'{dc}{r}'].font = BLACK
SENS_LAST = 5 + len(CASES) - 1
r = SENS_LAST + 3
M = json.load(open(MD.OUT, encoding='utf-8'))
ws.cell(row=r, column=1, value='회사 Y5 Contribution 민감도 (Base, 만원) — STATIC SNAPSHOT: model.py sens_company() 결과. 입력 변경 시 model.py 재실행 필요').font = BOLD
r += 1
hdr(ws, r, ['변수 (불리 ↔ 유리)', '불리 Δ', '유리 Δ'] + [''] * (len(PARAMS) - 2)); r += 1
ws.cell(row=r, column=1, value='Base Y5 Contribution').font = BOLD
c = ws.cell(row=r, column=2, value=round(M['sens_company']['base'], 1)); c.font = BLUE; c.number_format = NUM
r += 1
for d in M['sens_company']['items']:
    ws.cell(row=r, column=1, value=d['name']).font = BLACK
    for col, v in ((2, d['lo']), (3, d['hi'])):
        c = ws.cell(row=r, column=col, value=round(v, 1)); c.font = BLUE; c.number_format = NUM
    r += 1

# ---------------------------------------------------------------- Budget_24M (팀 · 사용처 · TIPS 과제 · 재원 · Seed)
ws = wb.create_sheet('Budget_24M')
setw(ws, [8, 40, 10, 9, 13, 9, 10, 9, 10, 10, 13, 13, 13, 13])
F = M['funding']; TPj = F['tips']
ws['A1'] = '24개월 자금 계획 (만원) — 팀 계획 · 사용처 · TIPS 과제 편성 · 재원 · Seed 범위'; ws['A1'].font = TITLE
ws['A2'] = ('팀 · 지출 표의 파란 값이 Inputs_Yearly Y1~Y2 지출의 원천 (model.py TEAM · COSTS). Founder 2인은 TIPS 요건 기준의 자리이며 인물 정보는 비어 있음. '
            '모든 값 ASSUMPTION, 정부지원 규정은 FACT (2026 공고, 선정 미확정).')
hdr(ws, 4, ['Key', '역할', '구분', '시작 월', '인당 연 인건비', 'FTE 비율', 'TIPS 참여율', 'Lean 포함',
            'Y1 인월', 'Y2 인월', 'Y1 인건비', 'Y2 인건비', 'TIPS 인건비 Y1', 'TIPS 인건비 Y2'])
r = 5; T0 = r
for key, role, m0, cost, frac, part, kind, lean, lean_m0 in MD.TEAM:
    ws.cell(row=r, column=1, value=key).font = BLACK
    ws.cell(row=r, column=2, value=role).font = BLACK
    ws.cell(row=r, column=3, value=M['team_kind'][kind]).font = BLACK
    for col, v, nf in (('D', m0, NUM), ('E', cost, NUM), ('F', frac, NUM1), ('G', part, PCT)):
        c = ws[f'{col}{r}']; c.value = v; c.font = BLUE; c.number_format = nf
    ws.cell(row=r, column=8, value=('예' if lean else '아니오') + (f' (M{lean_m0} 시작)' if lean_m0 else '')).font = BLACK
    ws[f'I{r}'] = f"=IF(D{r}<=12,13-D{r},0)*F{r}"
    ws[f'J{r}'] = f"=IF(D{r}<=13,12,MAX(0,25-D{r}))*F{r}"
    ws[f'K{r}'] = f"=E{r}*I{r}/12"; ws[f'L{r}'] = f"=E{r}*J{r}/12"
    ws[f'M{r}'] = f"=IF(OR(C{r}=\"R&D\",C{r}=\"Founder\"),K{r}*G{r},0)"
    ws[f'N{r}'] = f"=IF(OR(C{r}=\"R&D\",C{r}=\"Founder\"),L{r}*G{r},0)"
    for col in 'IJKLMN':
        ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM1
    r += 1
T1 = r - 1
ws.cell(row=r, column=2, value='합계').font = BOLD
for col in 'IJKLMN':
    ws[f'{col}{r}'] = f"=SUM({col}{T0}:{col}{T1})"; ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM1
TEAM_TOT = r; r += 1
ws.cell(row=r, column=2, value='평균 FTE (인월 ÷ 12)').font = BLACK
for col in 'IJ':
    ws[f'{col}{r}'] = f"={col}{TEAM_TOT}/12"; ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = '0.00'
FTE_R = r; r += 1
ws.cell(row=r, column=2, value='검증: FM_Base 인건비 − 팀 계획 (0이어야 함)').font = BLACK
ws[f'K{r}'] = f"={fmref('B', 'op_people', 0)}-K{TEAM_TOT}"; ws[f'L{r}'] = f"={fmref('B', 'op_people', 1)}-L{TEAM_TOT}"
for col in 'KL': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM1
PEOPLE_CHK = r; r += 2

hdr(ws, r, ['', 'B. 인건비 외 사용처 (Spec 33 항목)', '', '', 'Y1', 'Y2', '합계', 'TIPS 비율 Y1', 'TIPS 비율 Y2',
            'TIPS 편성 Y1', 'TIPS 편성 Y2', 'TIPS 비목', '', '']); r += 1
C0 = r
for key, cat, item, by, ly, sh, tcat in MD.COSTS:
    ws.cell(row=r, column=1, value=key).font = BLACK
    ws.cell(row=r, column=2, value=f'{cat} — {item}').font = BLACK
    for col, v, nf in (('E', by[0], NUM), ('F', by[1], NUM), ('H', sh[0], PCT), ('I', sh[1], PCT)):
        c = ws[f'{col}{r}']; c.value = v; c.font = BLUE; c.number_format = nf
    ws[f'G{r}'] = f"=E{r}+F{r}"; ws[f'J{r}'] = f"=E{r}*H{r}"; ws[f'K{r}'] = f"=F{r}*I{r}"
    for col in 'GJK': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
    ws.cell(row=r, column=12, value=tcat).font = BLACK
    r += 1
C1 = r - 1
ws.cell(row=r, column=2, value='인건비 외 합계').font = BOLD
for col in 'EFGJK':
    ws[f'{col}{r}'] = f"=SUM({col}{C0}:{col}{C1})"; ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM
OTHER_TOT = r; r += 1
ws.cell(row=r, column=2, value='예비비 (인건비 외 × 10%)').font = BLACK
ws[f'E{r}'] = f"=E{OTHER_TOT}*{MD.CONTINGENCY}"; ws[f'F{r}'] = f"=F{OTHER_TOT}*{MD.CONTINGENCY}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
CONT_R = r; r += 1
ws.cell(row=r, column=2, value='연구수당 (TIPS 현금 인건비 × 5%)').font = BLACK
cash_pay = lambda col: f"SUMIF($C${T0}:$C${T1},\"R&D\",{col}${T0}:{col}${T1})"
ws[f'E{r}'] = f"={cash_pay('M')}*{MD.ALLOW_RATE}"; ws[f'F{r}'] = f"={cash_pay('N')}*{MD.ALLOW_RATE}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
ALLOW_R = r; r += 1
ws.cell(row=r, column=2, value='검증: FM_Base 기타 운영비 − (인건비 외 + 예비비 + 연구수당) (0이어야 함)').font = BLACK
ws[f'E{r}'] = f"={fmref('B', 'op_other', 0)}-(E{OTHER_TOT}+E{CONT_R}+E{ALLOW_R})"
ws[f'F{r}'] = f"={fmref('B', 'op_other', 1)}-(F{OTHER_TOT}+F{CONT_R}+F{ALLOW_R})"
for col in 'EF': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM1
OTHER_CHK = r; r += 2

hdr(ws, r, ['', 'C. TIPS 과제 편성 (비목)', '', '', '1차년도', '2차년도', '합계', '', '', '', '', '근거', '', '']); r += 1
P0 = r
def sum_cat(col, cat):
    return f"SUMIF($L${C0}:$L${C1},\"{cat}\",{col}${C0}:{col}${C1})"
PRJ = [('인건비 (현금)', f"=SUMIF($C${T0}:$C${T1},\"R&D\",M${T0}:M${T1})", f"=SUMIF($C${T0}:$C${T1},\"R&D\",N${T0}:N${T1})", 'R&D 인원 × 참여율'),
       ('인건비 (현물)', f"=SUMIF($C${T0}:$C${T1},\"Founder\",M${T0}:M${T1})", f"=SUMIF($C${T0}:$C${T1},\"Founder\",N${T0}:N${T1})", 'Founder 2인 × 참여율'),
       ('연구재료비', '=' + sum_cat('J', '연구재료비'), '=' + sum_cat('K', '연구재료비'), 'Robot HW · Hand · Mock-up × 편성 비율'),
       ('연구활동비', '=' + sum_cat('J', '연구활동비'), '=' + sum_cat('K', '연구활동비'), 'SW · Data · 인증 사전시험 · IP × 편성 비율'),
       ('연구수당', f"=E{ALLOW_R}", f"=F{ALLOW_R}", '현금 인건비 × 5%')]
for lab, f1, f2, note in PRJ:
    ws.cell(row=r, column=2, value=lab).font = BLACK
    ws[f'E{r}'] = f1; ws[f'F{r}'] = f2; ws[f'G{r}'] = f"=E{r}+F{r}"
    for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
    ws.cell(row=r, column=12, value=note).font = BLACK
    r += 1
DIR_R = r
ws.cell(row=r, column=2, value='직접비 합계').font = BOLD
for col in 'EFG':
    ws[f'{col}{r}'] = f"=SUM({col}{P0}:{col}{r - 1})"; ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM
r += 1
TOT_R = r + 1
ws.cell(row=r, column=2, value='간접비 (연차 총액 − 직접비, 직접비 비례 배분)').font = BLACK
ws[f'E{r}'] = f"=G{TOT_R}*E{DIR_R}/G{DIR_R}-E{DIR_R}"; ws[f'F{r}'] = f"=G{TOT_R}*F{DIR_R}/G{DIR_R}-F{DIR_R}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
IND_R = r; r += 1
ws.cell(row=r, column=2, value='TIPS 과제 총 연구개발비 (= 정부지원금 ÷ 정부 비율 상한)').font = BOLD
ws[f'G{r}'] = f"={b('tips')}/{b('tips_gov_ratio')}"; ws[f'E{r}'] = f"=E{DIR_R}+E{IND_R}"; ws[f'F{r}'] = f"=F{DIR_R}+F{IND_R}"
for col in 'EFG': ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM
assert r == TOT_R; r += 1
CHK = [('정부지원금 (선정 시)', f"={b('tips')}", NUM, 'FACT (2026 공고 최대)'),
       ('기관부담 (= 총액 − 정부지원금)', f"=G{TOT_R}-{b('tips')}", NUM, ''),
       ('기관부담 중 현물 (Founder 인건비)', f"=G{P0 + 1}", NUM, ''),
       ('기관부담 중 현금', f"=G{TOT_R}-{b('tips')}-G{P0 + 1}", NUM, 'Seed 자금으로 부담'),
       ('현금 비율 하한 충족 (1 = 예)', f"=IF((G{TOT_R}-{b('tips')}-G{P0 + 1})/(G{TOT_R}-{b('tips')})>={b('tips_cash_ratio')}-0.0001,1,0)", NUM, '기관부담 중 현금 ≥ tips_cash_ratio'),
       ('간접비율 (간접비 ÷ 현금 직접비)', f"=G{IND_R}/(G{DIR_R}-G{P0 + 1})", PCT, '협약 기준 확인 필요 (ASSUMPTION)')]
CHK_R = {}
for lab, fs, nf, note in CHK:
    ws.cell(row=r, column=2, value=lab).font = BLACK
    c = ws[f'G{r}']; c.value = fs; c.number_format = nf; c.font = BLACK
    ws.cell(row=r, column=12, value=note).font = BLACK
    CHK_R[lab] = r; r += 1
r += 1

hdr(ws, r, ['', 'D. 재원 · Seed 범위 (24개월)', '', '', 'Y1', 'Y2', '합계', '', '', '', '', '비고', '', '']); r += 1
SP_R = r
ws.cell(row=r, column=2, value='회사 전체 지출 (= −FM_Base 현금흐름)').font = BOLD
ws[f'E{r}'] = f"=-{fmref('B', 'cash', 0)}"; ws[f'F{r}'] = f"=-{fmref('B', 'cash', 1)}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BOLD; ws[f'{col}{r}'].number_format = NUM
r += 1
ws.cell(row=r, column=2, value='  그중 실증 순비용 (지출 − Opex)').font = BLACK
ws[f'E{r}'] = f"=E{SP_R}-{fmref('B', 'opex', 0)}"; ws[f'F{r}'] = f"=F{SP_R}-{fmref('B', 'opex', 1)}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
r += 1
ws.cell(row=r, column=2, value='TIPS 정부지원금 (선정 시, 연차 배분)').font = BLACK
ws[f'E{r}'] = f"=E{TOT_R}*{b('tips_gov_ratio')}"; ws[f'F{r}'] = f"=F{TOT_R}*{b('tips_gov_ratio')}"; ws[f'G{r}'] = f"=E{r}+F{r}"
for col in 'EFG': ws[f'{col}{r}'].font = BLACK; ws[f'{col}{r}'].number_format = NUM
GOV_R = r; r += 1
ws.cell(row=r, column=2, value=f'Buffer (Y2 월평균 지출 × {MD.BUFFER_MONTHS}개월, Series A 협상 기간)').font = BLACK
ws[f'G{r}'] = f"=F{SP_R}/12*{MD.BUFFER_MONTHS}"; ws[f'G{r}'].font = BLACK; ws[f'G{r}'].number_format = NUM
BUF_R = r; r += 1
ws.cell(row=r, column=2, value='Seed 필요액 — Base (지출 − 정부지원금 + Buffer)').font = BOLD
ws[f'G{r}'] = f"=G{SP_R}-G{GOV_R}+G{BUF_R}"; ws[f'G{r}'].font = BOLD; ws[f'G{r}'].number_format = NUM; ws[f'G{r}'].fill = KEY_FILL
SEED_R = r; r += 1
for lab, v, note in [('Seed 필요액 — Lean (팀 축소 · 목업/실증/인증 축소) — STATIC', F['seed_lean'], 'model.py budget24(\'lean\')'),
                     ('Seed 필요액 — TIPS 미선정 시 Lean 범위 단독 — STATIC', F['seed_no_tips'], '연구수당 제외 · 정부지원 0'),
                     ('TIPS 운영사 선투자 요건 (수도권, FACT)', MD.V('B')['op_invest_min'], 'Seed 라운드의 일부로 충족')]:
    ws.cell(row=r, column=2, value=lab).font = BLACK
    c = ws[f'G{r}']; c.value = round(v, 1); c.font = BLUE; c.number_format = NUM
    ws.cell(row=r, column=12, value=note).font = BLACK
    r += 1
BUD_ROWS = dict(team_tot=TEAM_TOT, fte=FTE_R, people_chk=PEOPLE_CHK, other_tot=OTHER_TOT, other_chk=OTHER_CHK,
                cont=CONT_R, allow=ALLOW_R, dir=DIR_R, ind=IND_R, tot=TOT_R, spend=SP_R, gov=GOV_R, buf=BUF_R, seed=SEED_R,
                **{k: v for k, v in CHK_R.items()})

# ---------------------------------------------------------------- Sources
ws = wb.create_sheet('Sources')
setw(ws, [44, 18, 14, 90])
ws['A1'] = 'FACT Sources (조회일 2026-10-07~08, 검색 결과 기준 — 외부 제출 전 원문 재확인 필요)'; ws['A1'].font = TITLE
hdr(ws, 3, ['항목', '값', '기준', 'Source'])
src = json.load(open(os.path.join(HERE, 'sources.json'), encoding='utf-8'))
for i, d in enumerate(src):
    for j, k in enumerate(('item', 'value', 'basis', 'source')):
        c = ws.cell(row=4 + i, column=1 + j, value=d[k]); c.font = BLACK
        c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save(OUT)
print('saved', OUT)
json.dump({'fm_rows': FM_ROW, 'hh_rows': hrow, 'ue_rows': ue_row, 'ue_extra': UE_EXTRA, 'mk_rows': MK_ROW, 'bud_rows': BUD_ROWS}, open(os.path.join(HERE, '_xlsx_rows.json'), 'w'))
