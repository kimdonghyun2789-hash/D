# Helper kit reproducing the original SoftHand IR deck components (python-pptx).
import copy, re
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree
from PIL import ImageFont

# ---------------------------------------------------------------- tokens
BG = '0B0C0E'; CARD = '15171B'; CARD2 = '1D2025'; CARD3 = '25292F'
LINE = '2E333B'; LINE2 = '3C424B'
TEXT = 'F2F2EF'; TEXT2 = 'AEB3BB'; MUTED = '8A8F97'; MUTED2 = '737882'; DIM = '51565E'
OR = 'EC7A3C'; OR_FILL = '2E1D14'; OR_LIGHT = 'F3A574'
YEL = 'E8B649'; YEL_FILL = '2C2615'
GREEN = '5FB98E'
# light (appendix)
L_BG = 'F5F4F1'; L_TEXT = '16181B'; L_TEXT2 = '51565E'; L_LINE = 'DCDAD5'; L_CARD = 'FFFFFF'; L_MUTED = '8A8F97'

FONTS = {'R': 'Noto Sans KR', 'M': 'Noto Sans KR Medium', 'S': 'Noto Sans KR SemiBold', 'B': 'Noto Sans KR Bold', 'K': 'Noto Sans KR Black'}
# Noto Sans KR static TTFs (weights 400/500/600/700/900) are used only for the text-fit check.
# Put them in ~/.fonts (or set NOTO_KR_DIR); without them the fit check falls back to an approximation.
import os as _os
_WEIGHT = {'R': 400, 'M': 500, 'S': 600, 'B': 700, 'K': 900}
_DIRS = [_os.environ.get('NOTO_KR_DIR', ''), _os.path.expanduser('~/.fonts'), _os.path.expanduser('~/.local/share/fonts')]
def _ttf(f):
    for d in _DIRS:
        p = _os.path.join(d, f'NotoSansKR-{_WEIGHT[f]}.ttf')
        if d and _os.path.exists(p): return p
    return None
_FC = {}
def _font(f):
    if f not in _FC:
        p = _ttf(f); _FC[f] = ImageFont.truetype(p, 200) if p else None
    return _FC[f]
def text_width_in(t, f='R', pt=12, spc=0):
    fo = _font(f)
    if fo is None:  # approximation: Hangul ~0.92em, other ~0.55em
        em = sum(0.92 if '\uac00' <= ch <= '\ud7a3' else 0.55 for ch in t)
        return em * pt / 72 + max(0, len(t) - 1) * spc / 100 / 72
    return fo.getlength(t) / 200 * pt / 72 + max(0, len(t) - 1) * spc / 100 / 72

FIT_LOG = []   # (slide_label, name, problem)
CUR = {'label': ''}

def I(v): return Emu(int(round(v * 914400)))

# ---------------------------------------------------------------- presentation
def open_base(src):
    prs = Presentation(src)
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId); lst.remove(sid)
    return prs

def new_slide(prs, label, bg=BG):
    layout = [l for l in prs.slide_layouts if l.name == 'Blank'][0]
    s = prs.slides.add_slide(layout)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    CUR['label'] = label
    return s

def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt

# ---------------------------------------------------------------- text
def _set_run_style(r, f, size, color, spc=0, bold=False):
    rPr = r._r.get_or_add_rPr()
    rPr.set('sz', str(int(round(size * 100)))); rPr.set('b', '1' if bold else '0'); rPr.set('i', '0')
    if spc: rPr.set('spc', str(int(spc)))
    for tag in ('a:solidFill', 'a:latin', 'a:ea', 'a:cs'):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
    sf = etree.SubElement(rPr, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        e = etree.SubElement(rPr, qn(tag)); e.set('typeface', FONTS[f])

def _norm_paras(paras, d):
    """paras: str | list of (str | list[run] | dict). run: str or (text, {style})"""
    if isinstance(paras, str): paras = paras.split('\n')
    out = []
    for p in paras:
        if isinstance(p, dict):
            runs = p.get('runs', [p.get('text', '')]); st = {k: v for k, v in p.items() if k not in ('runs', 'text')}
        else:
            runs = p if isinstance(p, list) else [p]; st = {}
        rr = []
        for r in runs:
            if isinstance(r, tuple): rr.append((r[0], {**d, **st, **r[1]}))
            else: rr.append((r, {**d, **st}))
        out.append((rr, {**d, **st}))
    return out

def text(slide, x, y, w, h, paras, size=12, f='R', color=TEXT, align='l', anchor='t', ls=1.12, spc=0, after=0,
         name=None, bold=False, check=True, wrap=True, inset=0.0):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = wrap; tf.auto_size = MSO_AUTO_SIZE.NONE
    bp = tf._txBody.bodyPr
    for k in ('lIns', 'rIns', 'tIns', 'bIns'): bp.set(k, str(int(inset * 914400)))
    bp.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])
    d = dict(size=size, f=f, color=color, spc=spc, bold=bold, align=align, ls=ls, after=after)
    P = _norm_paras(paras, d)
    for i, (runs, st) in enumerate(P):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[st['align']]
        p.line_spacing = st['ls']
        if st.get('after'): p.space_after = Pt(st['after'])
        if st.get('before'): p.space_before = Pt(st['before'])
        for t, rs in runs:
            r = p.add_run(); r.text = t
            _set_run_style(r, rs['f'], rs['size'], rs['color'], rs.get('spc', 0), rs.get('bold', False))
    if check: _check_fit(name or 'text', P, w - 2 * inset, h - 2 * inset, wrap)
    return tb

def _wrap_lines(runs, width_in):
    # word-wrap simulation over runs (tokens keep run style); returns line count
    toks = []
    for t, rs in runs:
        for part in re.split(r'(\s+)', t):
            if part: toks.append((part, rs))
    lines = 1; cur = 0.0; maxword = 0.0
    for part, rs in toks:
        wdt = text_width_in(part, rs['f'], rs['size'], rs.get('spc', 0))
        if part.isspace():
            cur += wdt; continue
        maxword = max(maxword, wdt)
        if cur + wdt > width_in + 1e-3 and cur > 0:
            lines += 1; cur = wdt
            # very long token that itself exceeds the width -> char wrap
            while cur > width_in: lines += 1; cur -= width_in
        else:
            cur += wdt
    return lines, maxword

def _check_fit(name, P, w, h, wrap=True):
    total = 0.0
    for runs, st in P:
        size = max([rs['size'] for _, rs in runs] or [st['size']])
        txt = ''.join(t for t, _ in runs)
        if not txt.strip():
            total += size * 1.448 * st['ls'] / 72; continue
        n, maxword = _wrap_lines(runs, w) if wrap else (1, 0)
        if not wrap:
            tw = sum(text_width_in(t, rs['f'], rs['size'], rs.get('spc', 0)) for t, rs in runs)
            if tw > w + 0.02: FIT_LOG.append((CUR['label'], name, f'no-wrap width {tw:.2f}>{w:.2f}: {txt[:40]}'))
        if maxword > w + 0.02: FIT_LOG.append((CUR['label'], name, f'word wider than box {maxword:.2f}>{w:.2f}: {txt[:40]}'))
        total += n * size * 1.448 * st['ls'] / 72 + (st.get('after', 0) or 0) / 72
    if total > h + 0.04:
        FIT_LOG.append((CUR['label'], name, f'height {total:.2f}>{h:.2f}: {"".join(t for t,_ in P[0][0])[:50]}'))

# ---------------------------------------------------------------- shapes
def _no_effects(shp):
    spPr = shp._element.spPr
    if spPr.find(qn('a:effectLst')) is None:
        etree.SubElement(spPr, qn('a:effectLst'))

def box(slide, x, y, w, h, fill=CARD, line=None, lw=0.75, r=0.08, dash=False, name=None, shape='round', shadow=False):
    st = {'round': MSO_SHAPE.ROUNDED_RECTANGLE, 'rect': MSO_SHAPE.RECTANGLE, 'oval': MSO_SHAPE.OVAL, 'pill': MSO_SHAPE.ROUNDED_RECTANGLE}[shape]
    shp = slide.shapes.add_shape(st, I(x), I(y), I(w), I(h))
    if name: shp.name = name
    if shape in ('round', 'pill'):
        shp.adjustments[0] = 0.5 if shape == 'pill' else min(0.5, r / max(min(w, h), 1e-3))
    if fill: shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line); shp.line.width = Pt(lw)
        if dash: shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: shp.line.fill.background()
    if shadow:
        spPr = shp._element.spPr
        eff = etree.SubElement(spPr, qn('a:effectLst'))
        sh = etree.SubElement(eff, qn('a:outerShdw')); sh.set('blurRad', '76200'); sh.set('dist', '19050'); sh.set('dir', '5400000'); sh.set('algn', 't'); sh.set('rotWithShape', '0')
        c = etree.SubElement(sh, qn('a:srgbClr')); c.set('val', '000000'); a = etree.SubElement(c, qn('a:alpha')); a.set('val', '14000')
    else: _no_effects(shp)
    tf = shp.text_frame; tf.text = ''
    return shp

def box_text(shp, paras, size=10, f='S', color=TEXT, align='c', anchor='m', ls=1.05, inset=0.05, spc=0, check=True, name=None):
    tf = shp.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    bp = tf._txBody.bodyPr
    for k in ('lIns', 'rIns'): bp.set(k, str(int(inset * 914400)))
    for k in ('tIns', 'bIns'): bp.set(k, str(int(0.02 * 914400)))
    bp.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])
    d = dict(size=size, f=f, color=color, spc=spc, bold=False, align=align, ls=ls, after=0)
    P = _norm_paras(paras, d)
    for i, (runs, st) in enumerate(P):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[st['align']]
        p.line_spacing = st['ls']
        for t, rs in runs:
            r = p.add_run(); r.text = t; _set_run_style(r, rs['f'], rs['size'], rs['color'], rs.get('spc', 0))
    if check:
        w = shp.width / 914400 - 2 * inset; h = shp.height / 914400 - 0.04
        _check_fit(name or shp.name, P, w, h)
    return shp

def pill(slide, x, y, w, h, t, color=TEXT2, line=LINE2, fill=None, size=8.5, f='S', spc=60, name=None, dash=False, right=None):
    if w is None: w = text_width_in(t, f, size, spc) + 0.34
    if x is None: x = right - w
    shp = box(slide, x, y, w, h, fill=fill, line=line, lw=0.75, shape='pill', dash=dash, name=name or f'Tag {t[:20]}')
    box_text(shp, t, size=size, f=f, color=color, spc=spc, inset=0.04)
    return shp

def tri(slide, x, y, s=0.1, color=MUTED, rot=90):
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, I(x), I(y), I(s), I(s))
    shp.rotation = rot; shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(color); shp.line.fill.background(); _no_effects(shp)
    return shp

def line(slide, x1, y1, x2, y2, color=LINE2, w=1.0, dash=False, arrow=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    if name: c.name = name
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(w)
    if dash: c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    return c

def dot(slide, x, y, d=0.1, color=OR):
    return box(slide, x, y, d, d, fill=color, shape='oval', name='dot')

def num_badge(slide, x, y, n, d=0.32, fill=OR, color='FFFFFF', size=11):
    b = box(slide, x, y, d, d, fill=fill, shape='oval', name=f'badge {n}')
    box_text(b, str(n), size=size, f='K', color=color, inset=0.0)
    return b

def picture(slide, path, x, y, w=None, h=None, crop=None, name=None):
    kw = {}
    if w is not None: kw['width'] = I(w)
    if h is not None: kw['height'] = I(h)
    p = slide.shapes.add_picture(path, I(x), I(y), **kw)
    if crop:
        l, t, r, b = crop
        p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = l, t, r, b
    if name: p.name = name
    return p

def picture_fit(slide, path, x, y, w, h, crop=None, name=None, align='c'):
    """place image inside box (contain) preserving aspect; crop given as fractions before fit"""
    from PIL import Image
    im = Image.open(path); iw, ih = im.size
    if crop: iw = iw * (1 - crop[0] - crop[2]); ih = ih * (1 - crop[1] - crop[3])
    ar = iw / ih
    if w / h > ar: ph = h; pw = h * ar
    else: pw = w; ph = w / ar
    px = x + (w - pw) / 2 if align == 'c' else (x if align == 'l' else x + w - pw)
    py = y + (h - ph) / 2
    p = slide.shapes.add_picture(path, I(px), I(py), width=I(pw), height=I(ph))
    if crop: p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    if name: p.name = name
    return p, (px, py, pw, ph)

def picture_cover(slide, path, x, y, w, h, name=None, fx=0.5, fy=0.5):
    """fill the box (cover) by cropping the image; fx/fy = focus point"""
    from PIL import Image
    im = Image.open(path); iw, ih = im.size; ar = iw / ih; br = w / h
    if ar > br:  # image wider -> crop sides
        keep = br / ar; extra = 1 - keep; l = extra * fx; crop = (l, 0, extra - l, 0)
    else:
        keep = ar / br; extra = 1 - keep; t = extra * fy; crop = (0, t, 0, extra - t)
    p = slide.shapes.add_picture(path, I(x), I(y), width=I(w), height=I(h))
    p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    if name: p.name = name
    return p

# ---------------------------------------------------------------- composite components
def header(slide, num, kicker, title, sub_en=None, tag=None, tag_style='normal', light=False, title_size=28):
    dot(slide, 0.6, 0.495, 0.1, OR)
    k = f'{num}  ·  {kicker}' if num else kicker
    text(slide, 0.8, 0.42, 9.5, 0.26, k, size=10.5, f='S', color=OR, spc=160, name='Kicker')
    text(slide, 0.6, 0.74, 12.13, 0.6, title, size=title_size, f='K', color=(L_TEXT if light else TEXT), ls=1.0, name='Title')
    if sub_en:
        text(slide, 0.6, 1.36, 12.13, 0.32, sub_en, size=13, f='M', color=(L_TEXT2 if light else MUTED), spc=20, name='Subtitle EN')
    if tag:
        w = max(1.25, text_width_in(tag, 'S', 8.5, 60) + 0.36)
        if tag_style == 'placeholder': pill(slide, 12.73 - w, 0.40, w, 0.27, tag, color=YEL, line=YEL, fill=YEL_FILL, dash=True)
        else: pill(slide, 12.73 - w, 0.40, w, 0.27, tag)

def footer(slide, page, appendix=False, light=False, extra=None):
    col = L_MUTED if light else MUTED2
    t = '[회사명 입력 필요]  ·  ' + ('Appendix' if appendix else '창업 및 Seed 투자 제안서') + '  ·  대외비'
    if extra: t = extra
    text(slide, 0.6, 7.08, 8.5, 0.24, t, size=9.5, f='R', color=col, name='Footer')
    text(slide, 11.53, 7.08, 1.2, 0.24, page, size=9.5, f='S', color=col, align='r', name='Page number')

def placeholder(slide, x, y, w, h, t, size=9.5, align='l'):
    b = box(slide, x, y, w, h, fill=YEL_FILL, line=YEL, lw=0.75, dash=True, r=0.06, name='Placeholder')
    box_text(b, t, size=size, f='M', color=YEL, align=align, inset=0.1)
    return b

def card(slide, x, y, w, h, fill=CARD, line=None, r=0.08, lw=0.75, dash=False, name='Card', shadow=False):
    return box(slide, x, y, w, h, fill=fill, line=line, r=r, lw=lw, dash=dash, name=name, shadow=shadow)

def label(slide, x, y, w, t, color=OR, size=9.5, spc=120, h=0.22, f='S', align='l'):
    return text(slide, x, y, w, h, t, size=size, f=f, color=color, spc=spc, align=align, name='Label')

def chip(slide, x, y, w, h, t, fill=CARD2, color=TEXT, line=None, size=9.5, f='S', r=0.05, dash=False):
    b = box(slide, x, y, w, h, fill=fill, line=line, r=r, dash=dash, name=f'Chip {t[:16]}')
    box_text(b, t, size=size, f=f, color=color, inset=0.04)
    return b

# ---------------------------------------------------------------- tables (appendix style, native)
TABLE_STYLE = '{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}'
def table(slide, x, y, w, col_w, rows, row_h=None, head=True, size=10, head_size=10, light=True, name='Table',
          col_align=None, first_col_bold=True, highlight_rows=(), highlight_color=None, line_color=None, text_color=None, text2=None):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, I(x), I(y), I(w), I(0.3 * nr))
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None: sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = TABLE_STYLE
    tot = sum(col_w)
    for i, cw in enumerate(col_w): tbl.columns[i].width = I(w * cw / tot)
    lc = line_color or (L_LINE if light else LINE)
    t1 = text_color or (L_TEXT if light else TEXT); t2 = text2 or (L_TEXT2 if light else TEXT2)
    for ri, row in enumerate(rows):
        if row_h: tbl.rows[ri].height = I(row_h[ri] if isinstance(row_h, (list, tuple)) else row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            is_head = head and ri == 0
            sz = head_size if is_head else size
            fnt = 'S' if (is_head or (ci == 0 and first_col_bold)) else 'R'
            col = t1 if (is_head or (ci == 0 and first_col_bold)) else t2
            if ri in highlight_rows: col = highlight_color or OR; fnt = 'S'
            if isinstance(val, tuple): val, ov = val; fnt = ov.get('f', fnt); col = ov.get('color', col); sz = ov.get('size', sz)
            tf = cell.text_frame; tf.word_wrap = True
            paras = str(val).split('\n')
            for pi, ptxt in enumerate(paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                al = (col_align[ci] if col_align else 'l')
                p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[al]
                p.line_spacing = 1.05
                r = p.add_run(); r.text = ptxt; _set_run_style(r, fnt, sz, col, 40 if is_head else 0)
            tcPr = cell._tc.get_or_add_tcPr()
            for k, v in (('marL', 73152), ('marR', 54864), ('marT', 27432), ('marB', 27432)): tcPr.set(k, str(v))
            tcPr.set('anchor', 'ctr')
            for e in list(tcPr): tcPr.remove(e)
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(side)); ln.set('w', '9525')
                if side == 'a:lnB':
                    sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', lc)
                else: etree.SubElement(ln, qn('a:noFill'))
            etree.SubElement(tcPr, qn('a:noFill'))
    # rough height check per row
    for ri, row in enumerate(rows):
        rh = (row_h[ri] if isinstance(row_h, (list, tuple)) else row_h) if row_h else 0.3
        for ci, val in enumerate(row):
            v = val[0] if isinstance(val, tuple) else val
            cw = w * col_w[ci] / tot - 0.14
            sz = head_size if (head and ri == 0) else size
            n = sum(_wrap_lines([(p, {'f': 'S', 'size': sz})], cw)[0] for p in str(v).split('\n'))
            need = n * sz * 1.448 * 1.05 / 72 + 0.06
            if need > rh + 0.02: FIT_LOG.append((CUR['label'], name, f'row {ri} col {ci} needs {need:.2f}>{rh:.2f}: {str(v)[:30]}'))
    return gf
