# Minimal design kit for the v2 SoftHand IR deck (python-pptx).
# Principles: two font weights, few text sizes, flat colors, no pills/glows/outlined cards,
# orange only for the single most important number or word on a slide.
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageFont

FONT = 'Noto Sans KR'
LH = 1.448           # Noto Sans KR natural line height (ascent + descent) in em; PowerPoint "1.0" spacing = this
W, H = 13.333, 7.5
MX = 0.75            # side margin
CW = W - 2 * MX
THEMES = {
    'light': dict(bg='FFFFFF', text='15171A', text2='4A4F57', muted='8C9198', line='E2E4E7', soft='F4F5F6',
                  accent='E2571B', accent_soft='FBEDE6', img_bg=(238, 240, 242), dark='15171A', on_dark='F3F3F1',
                  on_dark2='A9AEB5', grey_bar='C9CDD2'),
    'dark': dict(bg='0F1012', text='F3F3F1', text2='B4B8BE', muted='80868D', line='2C2F34', soft='191B1E',
                 accent='EE7434', accent_soft='2A1C14', img_bg=(25, 27, 30), dark='191B1E', on_dark='F3F3F1',
                 on_dark2='A9AEB5', grey_bar='4A4F57'),
}
T = dict(THEMES['light'])
NAME = {'theme': 'light'}

def set_theme(name):
    T.clear(); T.update(THEMES[name]); NAME['theme'] = name

# ---------------------------------------------------------------- fit check (Noto Sans KR TTF metrics)
_FD = [os.environ.get('NOTO_KR_DIR', ''), os.path.expanduser('~/.fonts')]
_FC = {}
def _font(bold):
    k = 700 if bold else 400
    if k not in _FC:
        p = next((os.path.join(d, f'NotoSansKR-{k}.ttf') for d in _FD
                  if d and os.path.exists(os.path.join(d, f'NotoSansKR-{k}.ttf'))), None)
        _FC[k] = ImageFont.truetype(p, 200) if p else None
    return _FC[k]

def text_w(t, size, bold=False):
    f = _font(bold)
    if f is None: return len(t) * size * 0.9 / 72
    return f.getlength(t) / 200 * size / 72

def n_lines(t, size, width, bold=False):
    """Greedy wrap at spaces (PowerPoint wraps Korean at word boundaries by default)."""
    lines = 1; cur = ''
    for w_ in t.split(' '):
        cand = (cur + ' ' + w_) if cur else w_
        if text_w(cand, size, bold) <= width: cur = cand
        else: lines += 1; cur = w_
    return lines

FIT = []
CUR = {'slide': ''}

def _flag(msg):
    FIT.append(f'[{CUR["slide"]}] {msg}')

# ---------------------------------------------------------------- deck
def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    th = prs.slide_master.part.part_related_by(
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme')
    root = etree.fromstring(th.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for tag in ('majorFont', 'minorFont'):
        fe = root.find(f'.//a:fontScheme/a:{tag}', ns)
        fe.find('a:latin', ns).set('typeface', FONT)
        fe.find('a:ea', ns).set('typeface', FONT)
    th._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    return prs

def new_slide(prs, name='', bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(bg or T['bg'])
    CUR['slide'] = name
    return s

def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt

# ---------------------------------------------------------------- primitives
def _rgb(c): return RGBColor.from_string(c)

def _font_runs(r, sz, b, color):
    f = r.font; f.name = FONT; f.size = Pt(sz); f.bold = b
    f.color.rgb = _rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None: e = etree.SubElement(rPr, qn(tag))
        e.set('typeface', FONT)

def _bullet(para, char, indent, color):
    pPr = para._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent)))); pPr.set('indent', str(-int(Inches(indent))))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', color)
    bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', FONT)
    bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', char)

def text(s, x, y, w, h, paras, size=14, bold=False, color=None, align='l', anchor='t', line=1.0,
         space_after=0, label=None, check=True, bullet=None, indent=0.2, bullet_color=None):
    """paras: str ('\\n' = new paragraph) or list; a paragraph is a str or a list of (text, opts) runs.
    opts: size, bold, color.  bullet: e.g. '•' or '–' adds a hanging bullet to every paragraph."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str): paras = paras.split('\n')
    need = 0.0
    avail = w - (indent if bullet else 0)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        para.line_spacing = line
        if space_after and i < len(paras) - 1: para.space_after = Pt(space_after)
        runs = [(p, {})] if isinstance(p, str) else p
        plain = ''; big = 0; anyb = False
        for t, o in runs:
            r = para.add_run(); r.text = t; plain += t
            sz = o.get('size', size); b = o.get('bold', bold)
            big = max(big, sz); anyb = anyb or b
            _font_runs(r, sz, b, o.get('color', color or T['text']))
        if bullet: _bullet(para, bullet, indent, bullet_color or T['muted'])
        n = n_lines(plain, big, avail - 0.03, anyb) if plain.strip() else 1
        need += n * big * LH * line / 72 + ((space_after / 72) if space_after and i < len(paras) - 1 else 0)
        for wd in plain.split(' '):
            if wd and text_w(wd, big, anyb) > avail: _flag(f'{label or plain[:20]!r}: word "{wd}" wider than box')
    if check and need > h + 0.02:
        _flag(f'{label or str(paras[0])[:24]!r}: needs {need:.2f}in > box {h:.2f}in')
    return tb

def text_h(paras, size, w, line=1.0, bold=False, space_after=0):
    """Estimated height of a text block (same rules as text())."""
    if isinstance(paras, str): paras = paras.split('\n')
    tot = 0
    for i, p in enumerate(paras):
        if isinstance(p, list): p = ''.join(rt for rt, _ in p)   # paragraph given as styled runs
        tot += n_lines(p, size, w - 0.03, bold) * size * LH * line / 72
        if space_after and i < len(paras) - 1: tot += space_after / 72
    return tot

def _nostyle(sh):
    """Drop the theme style reference python-pptx attaches (LibreOffice draws its shadow otherwise)."""
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)

def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = _rgb(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    _nostyle(sh)
    return sh

def dot(s, cx, cy, d=0.1, fill=None):
    return rect(s, cx - d / 2, cy - d / 2, d, d, fill=fill or T['text'], shape=MSO_SHAPE.OVAL)

def hline(s, x, y, w, color=None, lw=0.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = _rgb(color or T['line']); c.line.width = Pt(lw)
    _nostyle(c)
    return c

def vline(s, x, y, h, color=None, lw=0.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    c.line.color.rgb = _rgb(color or T['line']); c.line.width = Pt(lw)
    _nostyle(c)
    return c

IMG_CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_img')
def image(s, path, x, y, w, h, focus=(0.5, 0.5), bg=None, zoom=1.0):
    """Cover-fit: crop the image to the box aspect around `focus` (0..1), alpha composited on bg."""
    os.makedirs(IMG_CACHE, exist_ok=True)
    im = Image.open(path).convert('RGBA')
    base = Image.new('RGBA', im.size, tuple(bg or T['img_bg']) + (255,)); base.alpha_composite(im); im = base.convert('RGB')
    iw, ih = im.size; ar = w / h
    cw, ch = (int(ih * ar), ih) if iw / ih > ar else (iw, int(iw / ar))
    cw, ch = int(cw / zoom), int(ch / zoom)
    cx = min(max(int(focus[0] * iw - cw / 2), 0), iw - cw); cy = min(max(int(focus[1] * ih - ch / 2), 0), ih - ch)
    im = im.crop((cx, cy, cx + cw, cy + ch))
    maxw = int(w * 220)
    if im.width > maxw: im = im.resize((maxw, int(maxw / ar)), Image.LANCZOS)
    key = f'{os.path.splitext(os.path.basename(path))[0]}_{NAME["theme"]}_{int(w*100)}x{int(h*100)}_{int(focus[0]*100)}_{int(focus[1]*100)}_{int(zoom*100)}.jpg'
    out = os.path.join(IMG_CACHE, key); im.save(out, quality=90)
    return s.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))

def cutout(s, path, x, y, w, h, align='c', valign='m'):
    """Contain-fit a transparent cut-out image (product shot) inside the box, no crop."""
    im = Image.open(path); ar = im.width / im.height
    pw, ph = (h * ar, h) if w / h > ar else (w, w / ar)
    px = {'c': x + (w - pw) / 2, 'l': x, 'r': x + w - pw}[align]
    py = {'m': y + (h - ph) / 2, 't': y, 'b': y + h - ph}[valign]
    return s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))

def _cell(cell):
    """Table cell -> (text, opts). A cell is str, (str, opts) or a list of (text, opts) runs (one paragraph)."""
    if isinstance(cell, list):
        return ''.join(rt for rt, _ in cell), {'bold': any(ro.get('bold') for _, ro in cell)} if all(ro.get('bold') for _, ro in cell) else {}
    if isinstance(cell, tuple): return cell
    return cell, {}

# ---------------------------------------------------------------- table (no style, horizontal rules only)
NO_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
def table(s, x, y, w, header, rows, col_w=None, size=11, header_size=None, pad=0.06, align=None,
          bold_first_col=False, label='table', max_h=None, header_color=None, anchor='t'):
    """header: list of str (or None); rows: list of lists; a cell is str or (str, opts) with opts
    bold / color / fill / size / align. Returns total height."""
    header_size = header_size or size
    all_rows = ([header] if header else []) + rows
    nc = len(all_rows[0]); col_w = col_w or [w / nc] * nc
    align = align or ['l'] * nc
    heights = []
    for ri, r in enumerate(all_rows):
        hm = 0
        for ci, cell in enumerate(r):
            t, o = _cell(cell)
            is_h = header and ri == 0
            sz = o.get('size', header_size if is_h else size)
            b = o.get('bold', is_h or (bold_first_col and ci == 0))
            lines = sum(n_lines(p, sz, col_w[ci] - 0.16 - 0.03, b) for p in str(t).split('\n'))
            hm = max(hm, lines * sz * LH / 72)
        heights.append(hm + 2 * pad + 0.01)
    tot = sum(heights)
    if max_h and tot > max_h + 0.02: _flag(f'{label}: table {tot:.2f}in > {max_h:.2f}in')
    gf = s.shapes.add_table(len(all_rows), nc, Inches(x), Inches(y), Inches(w), Inches(tot))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None: sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NO_STYLE
    for ci, cw in enumerate(col_w): tbl.columns[ci].width = Inches(cw)
    for ri, hh in enumerate(heights): tbl.rows[ri].height = Inches(hh)
    for ri, r in enumerate(all_rows):
        is_h = header and ri == 0
        for ci, cell in enumerate(r):
            t, o = _cell(cell)
            runs = cell if isinstance(cell, list) else None
            c = tbl.cell(ri, ci)
            c.margin_left = Inches(0.08); c.margin_right = Inches(0.08)
            c.margin_top = Inches(pad); c.margin_bottom = Inches(pad)
            c.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE}[anchor]
            tf = c.text_frame; tf.word_wrap = True
            csz = o.get('size', header_size if is_h else size)
            cb = o.get('bold', is_h or (bold_first_col and ci == 0))
            ccol = o.get('color', (header_color or T['text']) if is_h else T['text'])
            if runs is not None:   # one paragraph of styled runs
                para = tf.paragraphs[0]
                para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align[ci]]
                para.line_spacing = 1.0
                for rt, ro in runs:
                    rr = para.add_run(); rr.text = rt
                    _font_runs(rr, ro.get('size', csz), ro.get('bold', cb), ro.get('color', ccol))
            else:
                for pi, ptxt in enumerate(str(t).split('\n')):
                    para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                    para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[o.get('align', align[ci])]
                    para.line_spacing = 1.0
                    rr = para.add_run(); rr.text = ptxt
                    _font_runs(rr, csz, cb, ccol)
            tcPr = c._tc.get_or_add_tcPr()
            for e in list(tcPr): tcPr.remove(e)
            def ln(tag, wd, color):
                el = etree.SubElement(tcPr, qn(tag)); el.set('w', str(wd))
                if color:
                    sf = etree.SubElement(el, qn('a:solidFill')); cl = etree.SubElement(sf, qn('a:srgbClr')); cl.set('val', color)
                else:
                    etree.SubElement(el, qn('a:noFill'))
            ln('a:lnL', 0, None); ln('a:lnR', 0, None)
            if is_h: ln('a:lnT', 12700, T['text'])
            else: ln('a:lnT', 0, None)
            if is_h: ln('a:lnB', 9525, T['text'])
            else: ln('a:lnB', 9525, T['line'])
            if o.get('fill'):
                sf = etree.SubElement(tcPr, qn('a:solidFill')); cl = etree.SubElement(sf, qn('a:srgbClr')); cl.set('val', o['fill'])
            else:
                etree.SubElement(tcPr, qn('a:noFill'))
    return tot

# ---------------------------------------------------------------- charts (native)
def _chart_fonts(chart, size, color):
    cs = chart._chartSpace
    txPr = cs.find(qn('c:txPr'))
    if txPr is None:
        txPr = etree.SubElement(cs, qn('c:txPr'))
        etree.SubElement(txPr, qn('a:bodyPr')); etree.SubElement(txPr, qn('a:lstStyle'))
        p = etree.SubElement(txPr, qn('a:p')); pPr = etree.SubElement(p, qn('a:pPr')); etree.SubElement(pPr, qn('a:defRPr'))
        etree.SubElement(p, qn('a:endParaRPr')).set('lang', 'ko-KR')
        # c:txPr must come before c:externalData / c:printSettings / c:userShapes / c:extLst
        for tag in ('c:externalData', 'c:printSettings', 'c:userShapes', 'c:extLst'):
            nxt = cs.find(qn(tag))
            if nxt is not None: nxt.addprevious(txPr); break
    for d in cs.iter(qn('a:defRPr')):
        d.set('sz', str(int(size * 100)))
        for tag in ('a:solidFill', 'a:latin', 'a:ea', 'a:cs'):
            for e in d.findall(qn(tag)): d.remove(e)
        sf = etree.SubElement(d, qn('a:solidFill')); cl = etree.SubElement(sf, qn('a:srgbClr')); cl.set('val', color)
        for tag in ('a:latin', 'a:ea', 'a:cs'): etree.SubElement(d, qn(tag)).set('typeface', FONT)

def _plot_layout(chart, px, py, pw, ph):
    pa = chart._chartSpace.find('.//' + qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.Element(qn('c:layout')); pa.insert(0, lay)
    for e in list(lay): lay.remove(e)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for tag, val in [('c:layoutTarget', 'inner'), ('c:xMode', 'edge'), ('c:yMode', 'edge'),
                     ('c:x', px), ('c:y', py), ('c:w', pw), ('c:h', ph)]:
        etree.SubElement(ml, qn(tag)).set('val', str(val))

def column_chart(s, x, y, w, h, cats, series, colors, stacked=False, vmax=None, fmt='0.0',
                 label_colors=None, show_labels=True, gap=70, plot=(0.02, 0.04, 0.96, 0.84), size=11, legend=False,
                 own_cats=True):
    cd = CategoryChartData(); cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    kind = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False; ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
    pl = ch.plots[0]; pl.gap_width = gap
    if stacked: pl.overlap = 100
    va = ch.value_axis; va.has_major_gridlines = False; va.visible = False
    va.minimum_scale = 0
    if vmax: va.maximum_scale = vmax
    va.major_tick_mark = XL_TICK_MARK.NONE
    ca = ch.category_axis; ca.major_tick_mark = XL_TICK_MARK.NONE; ca.has_major_gridlines = False
    ca.format.line.color.rgb = _rgb(T['line'])
    if own_cats:
        # category labels as text boxes (LibreOffice adds Hangul/digit spacing inside chart text)
        ca.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        px, py, pw, ph = plot
        for i, c in enumerate(cats):
            ccx = x + w * (px + pw * (i + 0.5) / len(cats))
            text(s, ccx - 0.7, y + h * (py + ph) + 0.08, 1.4, 0.3, c, size=size, color=T['text2'], align='c', check=False)
    for i, ser in enumerate(pl.series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = _rgb(colors[i]); ser.format.line.fill.background()
        ser.invert_if_negative = False
    if show_labels:
        pl.has_data_labels = True
        dl = pl.data_labels; dl.number_format = fmt; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
        if label_colors:
            for i, ser in enumerate(pl.series):
                sdl = ser.data_labels
                sdl.number_format = fmt; sdl.number_format_is_linked = False
                sdl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
                sdl.font.size = Pt(size); sdl.font.color.rgb = _rgb(label_colors[i])
                sdl.show_value = True
    _plot_layout(ch, *plot)
    _chart_fonts(ch, size, T['text2'])
    return gf

def bar_chart(s, x, y, w, h, cats, vals, color, labels=None, gap=45, plot=(0.36, 0.0, 0.5, 1.0), size=11, vmax=None,
              colors=None, own_cats=True):
    """Horizontal bars, first category on top. labels: custom per-point label text."""
    cd = CategoryChartData(); cd.categories = cats; cd.add_series('v', vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart; ch.has_title = False; ch.has_legend = False
    pl = ch.plots[0]; pl.gap_width = gap
    va = ch.value_axis; va.visible = False; va.has_major_gridlines = False; va.minimum_scale = 0
    if vmax: va.maximum_scale = vmax
    ca = ch.category_axis; ca.reverse_order = True; ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.fill.background()
    if own_cats:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        px, py, pw, ph = plot
        n = len(cats); lh = size * LH / 72
        for i, c in enumerate(cats):
            ccy = y + h * (py + ph * (i + 0.5) / n)
            text(s, x, ccy - lh / 2, w * px - 0.1, lh + 0.02, c, size=size, color=T['text'], align='r', label='bar cat ' + c)
    ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = _rgb(color); ser.format.line.fill.background()
    ser.invert_if_negative = False
    if colors:
        for i, c in enumerate(colors):
            pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _rgb(c)
    pl.has_data_labels = True
    dl = pl.data_labels; dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.number_format = '0.0'; dl.number_format_is_linked = False
    if labels:
        for i, t in enumerate(labels):
            p = ser.points[i]; tf = p.data_label.text_frame; tf.text = t
            p.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
            for para in tf.paragraphs:
                for r in para.runs: _font_runs(r, size, False, T['text'])
    _plot_layout(ch, *plot)
    _chart_fonts(ch, size, T['text'])
    return gf

# ---------------------------------------------------------------- slide furniture
def header(s, section, title, sub=None, size=30):
    """Small section label + headline (+ one-line sub). Returns the y just below the block."""
    y = 0.6
    if section:
        text(s, MX, y, 8, 0.3, section, size=12, bold=True, color=T['accent'], label='section')
        y += 0.42
    nl = title.count('\n') + 1
    th = nl * size * LH * 0.95 / 72
    text(s, MX, y, CW, th + 0.02, title, size=size, bold=True, line=0.95, label='title:' + title[:12])
    y += th + 0.14
    if sub:
        text(s, MX, y, CW, 0.36, sub, size=16, color=T['text2'], label='sub:' + sub[:12])
        y += 0.36
    return y

def footer(s, page, left='SoftHand-4  |  Seed 투자 제안서', note=None):
    text(s, MX, H - 0.45, 6, 0.22, left, size=9, color=T['muted'], check=False)
    if note:
        nx = MX + 2.7; nw = W - MX - 0.6 - nx
        if n_lines(note, 9, nw - 0.03) > 1: _flag(f'footer note wraps: {note[:30]!r}')
        text(s, nx, H - 0.45, nw, 0.22, note, size=9, color=T['muted'], align='r', check=False)
    text(s, W - MX - 0.5, H - 0.45, 0.5, 0.22, f'{page:02d}' if isinstance(page, int) else str(page),
         size=9, color=T['muted'], align='r', check=False)

def footnote(s, txt, y=None, w=None):
    """Small grey note just above the footer."""
    hh = text_h(txt, 10, w or CW)
    yy = y if y is not None else H - 0.62 - hh
    text(s, MX, yy, w or CW, hh + 0.02, txt, size=10, color=T['muted'], label='footnote')


# ---------------------------------------------------------------- v3 diagram helpers
def arrow(s, x1, y1, x2, y2, color=None, lw=1.25, head='triangle'):
    """Straight connector with an arrowhead at (x2, y2)."""
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(color or T['muted']); c.line.width = Pt(lw)
    _nostyle(c)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', head); te.set('w', 'med'); te.set('len', 'med')
    return c

def chip(s, x, y, w, h, txt, fill=None, color=None, size=12, bold=True, align='c', line=None, label=None, pad=0.08):
    """Filled block with centered text (no rounded corners, no outline unless `line`)."""
    rect(s, x, y, w, h, fill=fill or T['soft'], line=line)
    text(s, x + pad, y, w - 2 * pad, h, txt, size=size, bold=bold, color=color or T['text'], align=align, anchor='m',
         label=label or ('chip ' + str(txt)[:14]))

def dashed_rect(s, x, y, w, h, color=None, lw=1.0):
    from pptx.enum.dml import MSO_LINE
    sh = rect(s, x, y, w, h, line=color or T['grey_bar'], lw=lw)
    sh.line.dash_style = MSO_LINE.DASH
    return sh
