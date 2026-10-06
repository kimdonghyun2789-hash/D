# -*- coding: utf-8 -*-
# python3 source/deck/make_report.py OUT.md OUT.docx
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_content import *
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_MD = sys.argv[1]; OUT_DOCX = sys.argv[2]
A = M['assumptions']

def base_rows():
    f = lambda v: f'{v:.1f}'.replace('-', '−')
    rows = [['항목 (억 원)', '1년', '2년', '3년', '4년', '5년']]
    rows.append(['유료 PoC / 통합 (건)'] + [f'{a} / {b}' for a, b in zip(A['poc_n'], A['int_n'])])
    rows.append(['신규 핸드 (대) 직판 / 파트너'] + [f'{a} / {b}' for a, b in zip(A['hand_direct'], A['hand_partner'])])
    rows.append(['매출: 유료 PoC·통합'] + [f(v) for v in B['poc_int']])
    rows.append(['매출: SoftHand 하드웨어'] + [f(v) for v in B['hw']])
    rows.append(['매출: Kitchen Skill'] + [f(v) for v in B['skill']])
    rows.append(['매출: Runtime·유지보수'] + [f(v) for v in B['runtime']])
    rows.append(['매출: 파트너 판매'] + [f(v) for v in B['partner']])
    rows.append(['총매출'] + [f(v) for v in B['rev']])
    rows.append(['매출총이익 (이익률)'] + [f'{f(g)} ({m*100:.0f}%)' for g, m in zip(B['gp'], B['gm'])])
    rows.append(['운영비'] + [f(v) for v in B['opex']])
    rows.append(['영업손익'] + [f(v) for v in B['op']])
    rows.append(['재사용 매출 비중'] + [f'{v*100:.0f}%' for v in B['reuse_share']])
    return rows

ASSUMP = ('가정 (2차와 동일): 유료 PoC 건당 5,000만 원 · 통합 프로젝트 4,000만 원 · 핸드 1,500만 원 (파트너 순매출 1,200만 원) · 핸드 원가 950만 → 750만 원 · '
          'Kitchen Skill 300만 원 (파트너 240만 원, 핸드당 1.0 → 1.6개) · Runtime 설치 핸드당 연 150만 원 · 매출총이익률: PoC·통합 40%, Skill 85%, Runtime 70% · '
          '운영비 1·2년차는 Seed 집행 계획과 동일, 3~5년차는 평균 인원 14·19·23명 기준 · 매출원은 분야 무관 가정 (주방·산업 비중은 PoC 후 재산정)')
FUND = [f"18개월 핵심 운영 {S['core18']:.1f}억 원 + 6개월 연장 {S['ext6']:.1f}억 원 (M18 점검 통과 시) + 예비비 {CONT:.1f}억 원 = 20.0억 원 (2차와 동일)",
        f"손익 연결: 운영비 1년차 {S['opex_y1']:.1f}억 + 2년차 {S['opex_y2']:.1f}억 = 집행 계획 − 예비비",
        '항목명만 주방 중심으로 바꾸고 금액은 그대로 유지']
AVG = sum(sc for _, sc, _ in SCORES) / len(SCORES)

# ---------------------------------------------------------------- Markdown
def md_table(rows):
    out = ['| ' + ' | '.join(rows[0]) + ' |', '|' + '|'.join(['---'] * len(rows[0])) + '|']
    for r in rows[1:]: out.append('| ' + ' | '.join(str(c).replace('|', '/') for c in r) + ' |')
    return '\n'.join(out)

md = [f'# {TITLE}', f'_{SUBTITLE}_', '']
md.append(md_table([['구분', '내용']] + [list(r) for r in META]))
md += ['', '## 0. 요약', ''] + [f'- {t}' for t in SUMMARY]
md += ['', '### 2차 → 3차 비교', '']
md.append(md_table([['항목', '2차 수정본', '3차 수정본']] + [list(r) for r in METRICS]))
md += ['', '### 투자자가 도달해야 할 판단과 답하는 장', '']
md.append(md_table([['판단', '답하는 장']] + [list(r) for r in JUDGMENT]))
md += ['', '## 1. 슬라이드별 수정 내역', '', '### 1.1 판정 요약 (2차 본문 16장)', '']
md.append(md_table([['2차 장', '제목', '판정', '3차 위치']] + [[c[0], c[1], c[2], c[3]] for c in MAIN_CHANGES]))
md += ['', '### 1.2 장별 상세: 기존 메시지 → 수정 후 메시지 → 이유 → 투자자 관점 개선점', '']
for c in MAIN_CHANGES:
    md += [f'#### 2차 {c[0]} {c[1]} → 3차 {c[3]} · {c[2]}', '']
    md.append(md_table([['구분', '내용'], ['기존 메시지', c[4]], ['수정 후 메시지', c[5]], ['수정 이유', c[6]], ['투자자 관점 개선점', c[7]]]))
    md.append('')
md += ['### 1.3 신규 장', '']
md.append(md_table([['3차 장', '역할', '제목', '내용']] + [list(r) for r in NEW_SLIDES]))
md += ['', '### 1.4 부록 변경', '']
md.append(md_table([['2차 부록', '판정', '3차 위치', '내용']] + [list(r) for r in APPX_CHANGES]))
md += ['', '## 2. 핵심 변경 8가지의 이유', '']
for i, (t, items) in enumerate(WHYS):
    md += [f'### 2.{i + 1} {t}', ''] + [f'- {x}' for x in items] + ['']
md += ['## 3. 3차 수정본 구성', '']
md.append(md_table([['장', '제목', '역할']] + [list(r) for r in SLIDES]))
md += ['', '## 4. VC 관점 자체 검토 (완성본 기준)', '']
md.append(md_table([['질문', '검토 결과', '판정', '검토 후 수정']] + [list(r) for r in REVIEW]))
md += ['', '## 5. 투자심의 평가 (10점 만점)', '']
md.append(md_table([['항목', '점수', '근거']] + [[a, str(b), c] for a, b, c in SCORES]))
md += ['', f'평균 {AVG:.1f}점. 비전·구조 항목은 높고, 실물·고객·팀 항목이 낮은 전형적인 Pre-proof Seed 상태', '',
       '### 현재 상태에서 투자를 보류한다면 가장 큰 이유 3개', ''] + [f'{i + 1}. **{a}**: {b}' for i, (a, b) in enumerate(HOLD)]
md += ['', '### 이 3가지가 확보되면 실제 투자 결정 가능성이 크게 높아지는 증거 3개', ''] + [f'{i + 1}. **{a}**: {b}' for i, (a, b) in enumerate(EVIDENCE)]
md += ['', '## 6. 재무와 Seed 집행 (수치는 2차와 동일)', '', '### 6.1 Base Case = 하방 (대규모 주방·OEM 매출 제외)', '']
md.append(md_table(base_rows()))
md += ['', ASSUMP, '', '### 6.2 민감도', '']
md.append(md_table([['시나리오', '결과']] + [list(r) for r in SENS]))
md += ['', '### 6.3 Seed 20억 원 자금 사용: 항목명 변경 (억 원)', '']
md.append(md_table([['2차 항목명', '3차 항목명', '금액']] + [list(r) for r in UOF_RELABEL]))
md += [''] + [f'- {t}' for t in FUND] + ['']
md += ['## 7. 출처 재확인 결과', '']
md.append(md_table([['항목', '확인 내용']] + [list(r) for r in FACTS]))
md += ['', '## 8. 사실·가정 구분 원칙', ''] + [f'- {t}' for t in PRINCIPLES]
md += ['', '## 9. 외부 제출 전 입력할 정보', '']
md.append(md_table([['항목', '위치', '내용']] + [list(r) for r in TODO]))
md += ['', '## 10. 이미지', '']
md.append(md_table([['파일 (assets/renders/kitchen)', '장면', '사용 위치']] + [list(r) for r in IMAGES]))
md += ['', IMAGE_NOTE, '', '## 11. 출처 (열람 2026-10-06)', '']
md.append(md_table([['번호', '출처', 'URL']] + [list(r) for r in SOURCES_FULL]))
md += ['', '## 12. 다음 단계', ''] + [f'{i + 1}. {t}' for i, t in enumerate(NEXT)] + ['']
open(OUT_MD, 'w', encoding='utf-8').write('\n'.join(md))

# ---------------------------------------------------------------- DOCX
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
for side in ('left_margin', 'right_margin'): setattr(sec, side, Cm(2.0))
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)
FONT_KO = '맑은 고딕'; FONT_EN = 'Malgun Gothic'
ORANGE = RGBColor(0xE2, 0x57, 0x1B); DARK = RGBColor(0x15, 0x17, 0x1A); GRAY = RGBColor(0x4A, 0x4F, 0x57)

def set_font(run, size=10, bold=False, color=DARK):
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    run.font.name = FONT_EN
    rPr = run._element.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), FONT_EN)
    rf.set(qn('w:eastAsia'), FONT_KO)

st = doc.styles['Normal']; st.font.name = FONT_EN; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_KO)
pf = st.paragraph_format; pf.space_after = Pt(4); pf.line_spacing = 1.25
# no automatic extra space between Hangul and Latin letters / digits (Word default is on)
_pPr = st.element.get_or_add_pPr()
_succ = ('w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
         'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr')
for _tag in ('w:autoSpaceDE', 'w:autoSpaceDN'):
    _e = OxmlElement(_tag); _e.set(qn('w:val'), '0')
    _nxt = next((_pPr.find(qn(t)) for t in _succ if _pPr.find(qn(t)) is not None), None)
    if _nxt is not None: _nxt.addprevious(_e)
    else: _pPr.append(_e)

def para(text, size=10, bold=False, color=DARK, after=4, align=None, before=0):
    p = doc.add_paragraph(); r = p.add_run(text); set_font(r, size, bold, color)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    return p

def heading(text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text); set_font(r, {1: 15, 2: 12, 3: 11}[level], True, ORANGE if level == 1 else DARK)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    return h

def bullet(text, style='List Bullet'):
    p = doc.add_paragraph(style=style); r = p.add_run(text); set_font(r, 10)
    p.paragraph_format.space_after = Pt(3)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor); tcPr.append(shd)

def table(rows, widths_cm, head=True, first_col_bold=False, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, wcm in enumerate(widths_cm): t.columns[i].width = Cm(wcm)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.width = Cm(widths_cm[ci])
            c.text = ''
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            hb = (head and ri == 0) or (first_col_bold and ci == 0)
            set_font(r, size, hb, DARK)
            if head and ri == 0: shade(c, 'E9EBEE')
    tblPr = t._tbl.tblPr; borders = OxmlElement('w:tblBorders')
    for b in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{b}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D5D8DC'); borders.append(e)
    anchor = None
    for tag in ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription'):
        anchor = tblPr.find(qn(tag))
        if anchor is not None: break
    if anchor is not None: anchor.addprevious(borders)
    else: tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

para(TITLE, 20, True, DARK, after=2)
para(SUBTITLE, 10.5, False, ORANGE, after=10)
table([['구분', '내용']] + [list(r) for r in META], [3.2, 13.8], first_col_bold=True)
heading('0. 요약')
for t_ in SUMMARY: bullet(t_)
heading('2차 → 3차 비교', 3)
table([['항목', '2차 수정본', '3차 수정본']] + [list(r) for r in METRICS], [4.2, 5.6, 7.2], first_col_bold=True)
heading('투자자가 도달해야 할 판단과 답하는 장', 3)
table([['판단', '답하는 장']] + [list(r) for r in JUDGMENT], [9.5, 7.5])
heading('1. 슬라이드별 수정 내역')
heading('1.1 판정 요약 (2차 본문 16장)', 2)
table([['2차 장', '제목', '판정', '3차 위치']] + [[c[0], c[1], c[2], c[3]] for c in MAIN_CHANGES], [1.6, 4.6, 5.4, 5.4], size=9)
heading('1.2 장별 상세: 기존 메시지 → 수정 후 메시지 → 이유 → 투자자 관점 개선점', 2)
for c in MAIN_CHANGES:
    heading(f'2차 {c[0]} {c[1]} → 3차 {c[3]} · {c[2]}', 3)
    table([['구분', '내용'], ['기존 메시지', c[4]], ['수정 후 메시지', c[5]], ['수정 이유', c[6]], ['투자자 관점 개선점', c[7]]], [3.4, 13.6], first_col_bold=True)
heading('1.3 신규 장', 2)
table([['3차 장', '역할', '제목', '내용']] + [list(r) for r in NEW_SLIDES], [1.4, 2.4, 5.4, 7.8], size=9)
heading('1.4 부록 변경', 2)
table([['2차 부록', '판정', '3차 위치', '내용']] + [list(r) for r in APPX_CHANGES], [4.2, 1.6, 1.6, 9.6], size=8.5)
heading('2. 핵심 변경 8가지의 이유')
for i, (t_, items) in enumerate(WHYS):
    heading(f'2.{i + 1}  {t_}', 2)
    for x in items: bullet(x)
heading('3. 3차 수정본 구성')
table([['장', '제목', '역할']] + [list(r) for r in SLIDES], [1.6, 9.0, 6.4], size=9)
heading('4. VC 관점 자체 검토 (완성본 기준)')
table([['질문', '검토 결과', '판정', '검토 후 수정']] + [list(r) for r in REVIEW], [3.6, 5.8, 2.0, 5.6], size=8.5)
heading('5. 투자심의 평가 (10점 만점)')
table([['항목', '점수', '근거']] + [[a, str(b), c] for a, b, c in SCORES], [3.6, 1.4, 12.0], first_col_bold=True, size=9)
para(f'평균 {AVG:.1f}점. 비전·구조 항목은 높고, 실물·고객·팀 항목이 낮은 전형적인 Pre-proof Seed 상태', 10, True, DARK, after=6)
heading('현재 상태에서 투자를 보류한다면 가장 큰 이유 3개', 3)
for i, (a, b) in enumerate(HOLD): para(f'{i + 1}. {a}: {b}', 10, after=3)
heading('이 3가지가 확보되면 실제 투자 결정 가능성이 크게 높아지는 증거 3개', 3)
for i, (a, b) in enumerate(EVIDENCE): para(f'{i + 1}. {a}: {b}', 10, after=3)
heading('6. 재무와 Seed 집행 (수치는 2차와 동일)')
heading('6.1 Base Case = 하방 (대규모 주방·OEM 매출 제외)', 2)
table(base_rows(), [5.0, 2.4, 2.4, 2.4, 2.4, 2.4], first_col_bold=True, size=8.5)
para(ASSUMP, 9, color=GRAY)
heading('6.2 민감도', 2)
table([['시나리오', '결과']] + [list(r) for r in SENS], [6.0, 11.0], first_col_bold=True)
heading('6.3 Seed 20억 원 자금 사용: 항목명 변경 (억 원)', 2)
table([['2차 항목명', '3차 항목명', '금액']] + [list(r) for r in UOF_RELABEL], [6.2, 7.8, 3.0], size=9)
for t_ in FUND: bullet(t_)
heading('7. 출처 재확인 결과')
table([['항목', '확인 내용']] + [list(r) for r in FACTS], [3.4, 13.6], first_col_bold=True, size=9)
heading('8. 사실·가정 구분 원칙')
for t_ in PRINCIPLES: bullet(t_)
heading('9. 외부 제출 전 입력할 정보')
table([['항목', '위치', '내용']] + [list(r) for r in TODO], [3.6, 2.0, 11.4], first_col_bold=True, size=9)
heading('10. 이미지')
table([['파일 (assets/renders/kitchen)', '장면', '사용 위치']] + [list(r) for r in IMAGES], [5.4, 8.8, 2.8], size=9)
para(IMAGE_NOTE, 9, color=GRAY)
heading('11. 출처 (열람 2026-10-06)')
table([['번호', '출처', 'URL']] + [list(r) for r in SOURCES_FULL], [1.2, 7.4, 8.4], size=8)
heading('12. 다음 단계')
for i, t_ in enumerate(NEXT): para(f'{i + 1}. {t_}', 10, after=3)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run('SoftHand IR Deck 수정 보고서 (3차)  ·  대외비  ·  '); set_font(r, 8, False, GRAY)
fld1 = OxmlElement('w:fldSimple'); fld1.set(qn('w:instr'), 'PAGE'); rr = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '1'; rr.append(t); fld1.append(rr); fp._p.append(fld1)
z = doc.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'), '100')
doc.save(OUT_DOCX)
print('ok')
