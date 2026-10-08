# Cross-check the recalculated xlsx against model.json (run after scripts/recalc.py).
import os, sys, json
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import model as MD
M = json.load(open(MD.OUT, encoding='utf-8'))
R = json.load(open(os.path.join(HERE, '_xlsx_rows.json')))
wb = load_workbook(os.path.join(HERE, '..', 'ARKI_Robotics_Financial_Model.xlsx'), data_only=True)
FCOLS = ['C', 'D', 'E', 'F', 'G']
bad = 0; n = 0

def cmp(label, x, y, tol=0.01):
    global bad, n
    n += 1
    if x is None or abs((x or 0) - y) > tol * max(1, abs(y)) * 0.01 + 1e-6:
        bad += 1; print('MISMATCH', label, x, y)

for s in MD.SC:
    ws = wb[f'FM_{MD.SCN[s]}']; L = M['scenarios'][s]
    for name, row in R['fm_rows'][s].items():
        if name in L and isinstance(L[name], list):
            for t in range(5):
                cmp(f'{s}.{name}.Y{t+1}', ws[f'{FCOLS[t]}{row}'].value, L[name][t])
ws = wb['Household']; hr = R['hh_rows']
for col, key in (('B', 'purchase_direct_Y3'), ('C', 'purchase_direct_Y5'), ('D', 'rental_direct_Y3'), ('E', 'rental_direct_Y5')):
    h = M['household'][key]
    cmp(key + '.rev5', ws[f"{col}{hr['5년 매출']}"].value, h['rev5'])
    cmp(key + '.contrib5', ws[f"{col}{hr['5년 Lifetime Contribution']}"].value, h['contrib5'])
    cmp(key + '.gp5', ws[f"{col}{hr['5년 매출총이익 (획득비용 전)']}"].value, h['gp5'])
ws = wb['Unit_Economics']; ur = R['ue_rows']
for col, lab in (('B', 'Y3'), ('C', 'Y5')):
    r = M['rental'][lab]
    cmp('rental.payback.' + lab, ws[f"{col}{ur['Payback (BOM ÷ (요금 − Care − Grip), 개월)']}"].value, r['payback'])
    cmp('rental.contrib.' + lab, ws[f"{col}{ur['월 Contribution']}"].value, r['contrib_m'])
    cmp('rental.bommax.' + lab, ws[f"{col}{ur['Payback 허들 충족 최대 BOM']}"].value, r['bom_max'])
ux = R['ue_extra']
cmp('partner_irr', ws[f"B{ux['irr_row']}"].value, M['partner_irr']['B']['irr_y'])
for col in 'BCD':
    cmp('bom_check_' + col, ws[f"{col}{ux['bom_check_row']}"].value + 1, 1)
ws = wb['Market']; mk = M['market']['B']
for row, key in ((22, 'tam'), (25, 'sam'), (26, 'som'), (13, 'fit'), (7, 'tri1'), (8, 'tri2')):
    cmp('market.' + key, ws[f'B{row}'].value, mk[key])
ws = wb['TIPS_Budget']; TR = R['tips_rows']; TP = M['tips']
cmp('tips.spend_total', ws[f"D{TR['spend']}"].value, TP['spend_total'])
cmp('tips.rnd_total', ws[f"D{TR['rnd']}"].value, TP['total'])
cmp('tips.src_total', ws[f"D{TR['src']}"].value, TP['src_total'])
cmp('tips.gov_ratio_ok', ws[f"D{TR['정부지원 비율 상한 충족 (1 = 예)']}"].value, 1)
cmp('tips.cash_ok', ws[f"D{TR['민간 현금 비율 하한 충족 (1 = 예)']}"].value, 1)
cmp('tips.private_cash', ws[f"D{TR['민간부담 중 현금']}"].value, TP['private_cash'])
ws = wb['Sensitivity']
cmp('sens.base', ws['M5'].value, M['sens_household']['base'])
print(f'checked {n} values, mismatches {bad}')
sys.exit(1 if bad else 0)
