# Schematic CONCEPT drawings: kitchen plans (Case A-D) and the robot reach section. Units in mm, drawn to scale.
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_PATTERN, MSO_LINE
from pptx.util import Pt
import kit
from kit import T, text, rect, hline, vline, alpha, dashed_rect, _rgb

ORANGE = 'E2571B'; ORANGE_SOFT = 'F6C9B3'; GREY_ZONE = 'C9CDD2'; DARK = '15171A'; COUNTER = 'E2E4E7'; TALL = 'B9BEC5'

def _nolog(fn, *a, **k):
    kit.NOLOG['on'] = True
    try: return fn(*a, **k)
    finally: kit.NOLOG['on'] = False

def _lbl(s, x, y, w, h, t, size=7, color=None, bold=False, align='c'):
    _nolog(text, s, x, y, w, h, t, size=size, bold=bold, color=color or T['text'], align=align, anchor='m', check=False)

CASES = {
    'A': dict(room=(2700, 2400), walls=[('top',), ('left',)],
              counters=[(0, 0, 2000, 600), (0, 600, 600, 1500)],
              items=[(600, 0, 800, 560, 'Sink', 'FFFFFF'), (0, 700, 560, 600, 'IH', 'FFFFFF'),
                     (0, 1400, 560, 700, 'Storage', 'FFFFFF'), (2000, 0, 700, 600, 'Tall\nDW·Home', TALL)],
              rail=[], home=(2150, 120, 400, 360), reach=(2350, 600, 1050),
              rzone=[(600, 0, 2100, 600), (1500, 600, 1200, 450)],
              hzone=[(600, 1050, 1500, 1350)], nogo=[(2100, 1500, 600, 900, 'Dining\nNo-go')]),
    'B': dict(room=(3300, 2600), walls=[('top',), ('bottom',)],
              counters=[(0, 0, 3300, 600), (0, 2000, 3300, 600)],
              items=[(600, 0, 900, 560, 'Sink', 'FFFFFF'), (1550, 0, 900, 560, 'Counter', COUNTER),
                     (2700, 0, 600, 600, 'Tall\nDW', TALL), (0, 0, 560, 600, 'Home', TALL),
                     (900, 2040, 900, 560, 'IH', 'FFFFFF'), (2000, 2040, 1200, 560, 'Storage', 'FFFFFF')],
              rail=[(250, 300, 3050, 300)], home=(80, 120, 400, 360), reach=None,
              rzone=[(0, 0, 3300, 950)], hzone=[(0, 950, 3300, 1050)], nogo=[(0, 2000, 800, 600, 'No-go')]),
    'C': dict(room=(3000, 2700), walls=[('top',), ('left',), ('right',)],
              counters=[(0, 0, 3000, 600), (0, 600, 600, 1800), (2400, 600, 600, 1500)],
              items=[(0, 0, 600, 600, 'Dock', TALL), (1000, 0, 900, 560, 'Sink', 'FFFFFF'),
                     (2400, 0, 600, 600, 'Tall\nDW', TALL), (40, 1000, 560, 700, 'IH', 'FFFFFF'),
                     (2440, 900, 560, 900, 'Storage', 'FFFFFF')],
              rail=[(600, 300, 2300, 300)], home=(100, 120, 400, 360), reach=(300, 300, 1200),
              rzone=[(0, 0, 3000, 900), (0, 600, 1000, 900)], hzone=[(1000, 900, 1400, 1800)],
              nogo=[(0, 2400, 900, 300, 'No-go'), (2400, 2100, 600, 600, 'No-go')]),
    'D': dict(room=(4200, 3000), walls=[('top',)],
              counters=[(0, 0, 4200, 600), (900, 1600, 2400, 900)],
              items=[(700, 0, 900, 560, 'Sink', 'FFFFFF'), (1700, 0, 1200, 560, 'Counter', COUNTER),
                     (3000, 0, 600, 600, 'Tall\nDW', TALL), (3600, 0, 600, 600, 'Dock', TALL),
                     (0, 0, 600, 600, 'Storage', 'FFFFFF'), (1500, 1700, 1200, 700, 'Island IH', 'FFFFFF')],
              rail=[(300, 300, 3700, 300)], home=(3700, 120, 400, 360), reach=None,
              rzone=[(0, 0, 4200, 950)], hzone=[(0, 950, 4200, 650)],
              nogo=[(900, 1600, 2400, 900, 'Island No-go'), (3500, 2500, 700, 500, 'Dining')]),
}

def plan(s, x, y, w, h, key):
    c = CASES[key]; rw, rd = c['room']
    k = min(w / rw, h / rd)
    ox = x + (w - rw * k) / 2; oy = y + (h - rd * k) / 2
    X = lambda v: ox + v * k; Yy = lambda v: oy + v * k
    rect(s, X(0), Yy(0), rw * k, rd * k, fill='FFFFFF', line=T['line'], lw=0.5)
    for zx, zy, zw, zh in c['hzone']:
        sh = rect(s, X(zx), Yy(zy), zw * k, zh * k, fill=GREY_ZONE); alpha(sh, 45)
    for zx, zy, zw, zh in c['rzone']:
        sh = rect(s, X(zx), Yy(zy), zw * k, zh * k, fill=ORANGE_SOFT); alpha(sh, 55)
    for cx, cy, cw, ch in c['counters']:
        sh = rect(s, X(cx), Yy(cy), cw * k, ch * k, fill=COUNTER); alpha(sh, 85)
    for ix, iy, iw, ih, lab, fill in c['items']:
        rect(s, X(ix), Yy(iy), iw * k, ih * k, fill=fill, line=T['muted'], lw=0.5)
        _lbl(s, X(ix), Yy(iy), iw * k, ih * k, lab, size=7)
    for nx, ny, nw, nh, lab in c['nogo']:
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, *[kit.Inches(v) for v in (X(nx), Yy(ny), nw * k, nh * k)])
        sh.fill.patterned(); sh.fill.pattern = MSO_PATTERN.WIDE_DOWNWARD_DIAGONAL
        sh.fill.fore_color.rgb = _rgb('8C9198'); sh.fill.back_color.rgb = _rgb('FFFFFF')
        sh.line.color.rgb = _rgb(DARK); sh.line.width = Pt(0.75); kit._nostyle(sh)
        _lbl(s, X(nx), Yy(ny), nw * k, nh * k, lab, size=7, bold=True)
    for x1, y1, x2, y2 in c['rail']:
        ln = hline(s, X(x1), Yy(y1), (x2 - x1) * k, color=ORANGE, lw=3.0)
        ln.line.dash_style = MSO_LINE.DASH
    for zx, zy, zw, zh in c['rzone']:          # reach boundary = Robot Working Zone outline (kept inside the room)
        sh = rect(s, X(zx) + 0.01, Yy(zy) + 0.01, zw * k - 0.02, zh * k - 0.02, line=ORANGE, lw=1.0)
        sh.line.dash_style = MSO_LINE.DASH
    hx, hy, hw, hh = c['home']
    rect(s, X(hx), Yy(hy), hw * k, hh * k, fill=ORANGE)
    for wall in c['walls']:
        side = wall[0]
        if side == 'top': hline(s, X(0), Yy(0), rw * k, color=DARK, lw=3)
        if side == 'bottom': hline(s, X(0), Yy(rd), rw * k, color=DARK, lw=3)
        if side == 'left': vline(s, X(0), Yy(0), rd * k, color=DARK, lw=3)
        if side == 'right': vline(s, X(rw), Yy(0), rd * k, color=DARK, lw=3)
    _lbl(s, X(0), Yy(rd) + 0.02, rw * k, 0.16, f'{rw:,} × {rd:,} mm (Kitchen 영역)', size=7, color=T['muted'])
    return k

def legend(s, x, y, size=8):
    items = [('Robot Home', 'home'), ('Rail (상부장 하단)', 'rail'), ('Reach 경계', 'reach'), ('Robot Working Zone', 'rz'),
             ('Human Working Zone', 'hz'), ('No-go Zone', 'ng')]
    cx = x
    for lab, kind in items:
        if kind == 'home': rect(s, cx, y + 0.04, 0.16, 0.12, fill=ORANGE)
        elif kind == 'rail':
            ln = hline(s, cx, y + 0.1, 0.22, color=ORANGE, lw=3); ln.line.dash_style = MSO_LINE.DASH
        elif kind == 'reach':
            sh = rect(s, cx, y + 0.03, 0.2, 0.14, line=ORANGE, lw=1.0); sh.line.dash_style = MSO_LINE.DASH
        elif kind == 'rz': alpha(rect(s, cx, y + 0.03, 0.2, 0.14, fill=ORANGE_SOFT), 55)
        elif kind == 'hz': alpha(rect(s, cx, y + 0.03, 0.2, 0.14, fill=GREY_ZONE), 45)
        elif kind == 'ng':
            sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, *[kit.Inches(v) for v in (cx, y + 0.03, 0.2, 0.14)])
            sh.fill.patterned(); sh.fill.pattern = MSO_PATTERN.WIDE_DOWNWARD_DIAGONAL
            sh.fill.fore_color.rgb = _rgb('8C9198'); sh.fill.back_color.rgb = _rgb('FFFFFF')
            sh.line.color.rgb = _rgb(DARK); sh.line.width = Pt(0.5); kit._nostyle(sh)
        tw = kit.text_w(lab, size) + 0.05
        text(s, cx + 0.27, y, tw + 0.05, 0.22, lab, size=size, color=T['text2'], anchor='m', check=False)
        cx += 0.27 + tw + 0.22
    return cx

# ---------------------------------------------------------------- section (elevation) drawing
def section(s, x, y, w, h, detail=True, dark=False):
    """Side section: base cabinet, counter 850, upper cabinet 1,450~2,250, rail, inverted arm, reach R700, human zone."""
    col = 'F3F3F1' if dark else DARK
    mut = 'A9AEB5' if dark else T['muted']
    room_w, room_h = 2600, 2400
    k = min(w / room_w, h / room_h)
    ox = x; oy = y + h                       # origin at floor-left (wall face)
    X = lambda v: ox + v * k; Yy = lambda v: oy - v * k
    hline(s, X(0), Yy(0), room_w * k, color=col, lw=1.5)               # floor
    vline(s, X(0), Yy(2350), 2350 * k, color=col, lw=2.5)               # wall
    hline(s, X(0), Yy(2350), room_w * k, color=mut, lw=0.75)             # ceiling
    rect(s, X(0), Yy(850), 600 * k, 850 * k, fill=('2C2F34' if dark else COUNTER))           # base cabinet
    rect(s, X(0), Yy(880), 620 * k, 30 * k, fill=col)                                         # counter top
    rect(s, X(0), Yy(2250), 350 * k, 800 * k, fill=('2C2F34' if dark else COUNTER))          # upper cabinet
    rect(s, X(0), Yy(1450), 360 * k, 40 * k, fill=ORANGE)                                     # rail
    # inverted arm: carriage -> shoulder -> elbow -> wrist (simple links)
    pts = [(180, 1410), (260, 1260), (520, 1080), (520, 930)]
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        c = s.shapes.add_connector(1, kit.Inches(X(x1)), kit.Inches(Yy(y1)), kit.Inches(X(x2)), kit.Inches(Yy(y2)))
        c.line.color.rgb = _rgb(ORANGE); c.line.width = Pt(4); kit._nostyle(c)
    for px, py in pts[1:]:
        rect(s, X(px) - 0.04, Yy(py) - 0.04, 0.08, 0.08, fill=col, shape=MSO_SHAPE.OVAL)
    # reach envelope R700 from shoulder
    sx, sy, rr = 260, 1260, 700
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, *[kit.Inches(v) for v in (X(sx - rr), Yy(sy + rr), 2 * rr * k, 2 * rr * k)])
    sh.fill.background(); sh.line.color.rgb = _rgb(ORANGE); sh.line.width = Pt(1); sh.line.dash_style = MSO_LINE.DASH
    kit._nostyle(sh)
    # human (aisle)
    hx = 1500
    rect(s, X(hx), Yy(1450), 280 * k, 1450 * k, fill=('4A4F57' if dark else 'C9CDD2'), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, X(hx + 40), Yy(1700), 200 * k, 220 * k, fill=('4A4F57' if dark else 'C9CDD2'), shape=MSO_SHAPE.OVAL)
    if detail:
        sz = 8
        _lbl(s, X(620) + 0.05, Yy(880) - 0.1, 0.9, 0.2, 'Counter 850', size=sz, color=mut, align='l')
        _lbl(s, X(360) + 0.05, Yy(1450) - 0.1, 1.4, 0.2, 'Rail 하단 약 1,410', size=sz, color=ORANGE, align='l')
        _lbl(s, X(360) + 0.05, Yy(2250) - 0.05, 1.5, 0.2, '상부장 1,450~2,250', size=sz, color=mut, align='l')
        _lbl(s, X(sx + rr) - 0.6, Yy(sy + rr) - 0.02, 1.0, 0.2, 'Reach R700 (가정)', size=sz, color=ORANGE, align='l')
        _lbl(s, X(hx) - 0.3, Yy(0) - 0.02 - 1450 * k - 0.45, 1.2, 0.2, 'Human Zone', size=sz, color=mut)
        _lbl(s, X(0), Yy(2350) - 0.2, 1.4, 0.18, '천장고 약 2,300~2,400', size=sz, color=mut, align='l')
    return k
